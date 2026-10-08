#!/usr/bin/env python3
"""Build an editable PowerPoint of the Week 7 Day 1 deck (native text, not screenshots)."""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    req = Path(__file__).resolve().parent / "requirements.txt"
    sys.stderr.write(
        "Missing python-pptx. From the repo root run:\n"
        f"  python3 -m pip install --user -r {req}\n"
    )
    raise

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[1]
LOGO = REPO / "docs" / "branding" / "epic-learn-logo.png"
OUT = ROOT / "pdf" / "week-07-day-01-slides.pptx"
FOOTER = "Epic Learn Institute of Higher Education - Go Programming Master Course"

INK = RGBColor(0x14, 0x20, 0x2B)
MUTED = RGBColor(0x4D, 0x5D, 0x6B)
TEAL = RGBColor(0x0B, 0x6E, 0x68)
TEAL_DARK = RGBColor(0x08, 0x4F, 0x4B)
NAVY = RGBColor(0x0A, 0x2A, 0x5C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF4, 0xEF, 0xE6)
CARD = RGBColor(0xFF, 0xFD, 0xF9)
LINE = RGBColor(0xD8, 0xCF, 0xC2)
CODE_BG = RGBColor(0x10, 0x21, 0x2C)
CODE_FG = RGBColor(0xE8, 0xF2, 0xEF)
GOLD = RGBColor(0xD9, 0xB8, 0x26)
TITLE_BG = RGBColor(0x14, 0x2A, 0x32)
CREAM_TEXT = RGBColor(0xEA, 0xD9, 0xA3)
LIVE = RGBColor(0xC2, 0x4E, 0x1D)

W, H = 13.333, 7.5
MX = 0.65
HEADER_H = 0.78
FOOTER_H = 0.42
CONTENT_TOP = 0.95
TOTAL = 26


def _font(run, name="Calibri", size=18, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_textbox(slide, l, t, w, h, text, size=18, bold=False, color=INK, font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _font(run, font, size, bold, color)
    return box


def add_bullets(slide, l, t, w, h, items, size=20, color=INK):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        p.level = 0
        run = p.add_run()
        run.text = "•  " + item
        _font(run, "Calibri", size, False, color)
    return box


def add_code(slide, l, t, w, h, text, size=15):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = CODE_BG
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.1)
    lines = text.strip("\n").split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line if line else " "
        _font(run, "Consolas", size, False, CODE_FG)
    return shape


def add_card(slide, l, t, w, h, title, body, title_size=16, body_size=13):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = LINE
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    _font(run, "Calibri", title_size, True, TEAL_DARK)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    run2 = p2.add_run()
    run2.text = body
    _font(run2, "Calibri", body_size, False, INK)
    return shape


def add_brand(slide, page, dark=False):
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(HEADER_H))
    _fill(header, WHITE)
    slide.shapes.add_picture(str(LOGO), Inches(0.35), Inches(0.12), height=Inches(0.55))
    add_textbox(
        slide, 10.2, 0.22, 2.8, 0.4,
        "WEEK 7  ·  DAY 1",
        size=12, bold=True, color=CREAM_TEXT if dark else TEAL_DARK, align=PP_ALIGN.RIGHT,
    )
    gold = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(HEADER_H), Inches(W - 0.7), Inches(0.02)
    )
    _fill(gold, GOLD)

    footer_bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(H - FOOTER_H), Inches(W), Inches(FOOTER_H)
    )
    _fill(footer_bg, WHITE)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(H - FOOTER_H), Inches(W - 0.7), Inches(0.015)
    )
    _fill(line, NAVY)
    add_textbox(
        slide, 0.5, H - FOOTER_H + 0.06, 10.5, 0.3,
        FOOTER, size=11, color=NAVY, align=PP_ALIGN.LEFT,
    )
    add_textbox(
        slide, 11.3, H - FOOTER_H + 0.06, 1.6, 0.3,
        f"{page} / {TOTAL}", size=11, color=MUTED, align=PP_ALIGN.RIGHT,
    )


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_kicker_title(slide, kicker, title, title_size=32, top=None):
    top = CONTENT_TOP if top is None else top
    add_textbox(slide, MX, top, 12, 0.32, kicker.upper(), size=13, bold=True, color=TEAL)
    add_textbox(slide, MX, top + 0.28, 12, 0.7, title, size=title_size, bold=True, color=INK)


def new_slide(prs, page, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if dark:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
        _fill(bg, TITLE_BG)
    else:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
        _fill(bg, CREAM)
    add_brand(slide, page, dark=dark)
    return slide


def build(dest: Path | None = None) -> Path:
    dest = dest or OUT
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    page = 0

    def P():
        nonlocal page
        page += 1
        return page

    s = new_slide(prs, P(), dark=True)
    add_textbox(s, MX, 2.15, 12, 0.4, "GO PROGRAMMING MASTER COURSE", size=14, bold=True, color=CREAM_TEXT)
    add_textbox(s, MX, 2.55, 12, 1.0, "Week 7, Day 1", size=48, bold=True, color=WHITE)
    add_textbox(s, MX, 3.7, 11, 1.1, "Structs, JSON tags, and in-memory students — create and list via Postman.", size=22, color=RGBColor(0xE7, 0xDD, 0xD0))
    add_textbox(s, MX, 5.3, 11, 0.9, "Umesh Indrajith · BSc Eng (Hons), University of Moratuwa\nSenior Software Engineer / Lecturer", size=16, color=RGBColor(0xD7, 0xCB, 0xB8))
    add_notes(s, "Product week 4. Handlers from week 6 must exist. Today: models.Student, map store, POST+GET students. No JWT. No Postgres.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Since last week", "Handlers extracted. Now give them data shapes.", 26)
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Health + hello still 200. internal/handlers is real.",
        "Cold-call: method vs function? Why pointer receiver on handlers?",
        "Today we add structs and an in-memory student store.",
        "Still no database. Still no JWT. Product grows in memory first.",
    ], 20)
    add_notes(s, "Fix broken week-6 refactors in first 10 minutes only.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Today", "If Day 1 of week 7 works, you can")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Define a struct with exported fields and composite literals.",
        'Round-trip JSON with tags like `json:"first_name"`.',
        "Explain a tiny interface (error) and accept interfaces, return structs.",
        "Store students in a map[int]Student (or slice).",
        "Ship POST /api/students, GET /api/students, GET /api/students/:id.",
    ], 18)
    add_notes(s, "Update/delete are homework. Lab is create + list.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Software engineering", "A model is a shared language", 28)
    add_textbox(s, MX, 2.1, 12, 1.1, "JSON field names are a contract with Postman and later the React client. Tags make Go names and wire names both sane.", size=18, color=MUTED)
    add_bullets(s, MX, 3.4, 12, 3.0, [
        "SMS fields: first/last name, email, phone, course, date of birth.",
        "In-memory now → same shapes against Postgres in week 14.",
        "Invalid JSON → 400, not a panic.",
    ], 18)
    add_notes(s, "Point at COURSE.md / this repo models.Student.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Structs", "Named fields, one type")
    add_code(s, MX, 2.05, 12, 3.0, """type Student struct {
    ID        int
    FirstName string
    LastName  string
    Email     string
    Course    string
}
s := Student{FirstName: "Ada", LastName: "Lovelace", Course: "SE"}""", 14)
    add_bullets(s, MX, 5.2, 12, 1.6, [
        "Capital field = exported (visible outside the package / in JSON by default).",
        "Composite literal: set the fields you care about; others get zero values.",
    ], 16)
    add_notes(s, "Live type. Connect to week 2 zeros.")

    s = new_slide(prs, P())
    add_kicker_title(s, "JSON tags", "Wire names that match the API")
    add_code(s, MX, 2.05, 12, 3.0, """type Student struct {
    ID        int    `json:"id"`
    FirstName string `json:"first_name"`
    LastName  string `json:"last_name"`
    Email     string `json:"email"`
    Course    string `json:"course"`
}""", 14)
    add_bullets(s, MX, 5.2, 12, 1.6, [
        "Without tags, JSON would say FirstName. We want first_name.",
        "Match this repo’s model tags. Ignore DB tags until week 14.",
    ], 16)
    add_notes(s, "Show encode/decode with encoding/json or Gin ShouldBindJSON.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Methods on structs", "Behaviour next to data")
    add_code(s, MX, 2.05, 12, 2.2, """func (s Student) FullName() string {
    return s.FirstName + " " + s.LastName
}
// or reuse handlers.FullName(s.FirstName, s.LastName)""", 15)
    add_bullets(s, MX, 4.5, 12, 2.2, [
        "Value receiver is fine for read-only helpers.",
        "Store / handler that mutates the map → pointer receiver on the handler type.",
    ], 18)
    add_notes(s, "Keep FullName from week 6; optional method on Student.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Embedding · light", "Composition, not inheritance")
    add_code(s, MX, 2.05, 12, 2.6, """type Timestamps struct {
    CreatedAt time.Time `json:"created_at"`
}
type Student struct {
    ID int `json:"id"`
    Timestamps // embeds fields
}""", 15)
    add_bullets(s, MX, 4.9, 12, 1.8, [
        "Go has no class hierarchy. Embed when you want fields/methods promoted.",
        "Optional today — do not force embedding in the lab.",
    ], 18)
    add_notes(s, "Pace valve. One slide. Skip live if create/list eats time.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Interfaces · light", "Small surfaces. error is an interface.", 26)
    add_code(s, MX, 2.05, 12, 2.2, """type error interface {
    Error() string
}
// later: type StudentStore interface { Create(...); List() ... }""", 15)
    add_bullets(s, MX, 4.5, 12, 2.2, [
        "Accept interfaces, return structs — depend on behaviour, hand back concrete values.",
        "Today you may still return concrete []Student. Plant the proverb.",
    ], 18)
    add_notes(s, "Repository interface lands week 8/13. Do not over-abstract today.")

    s = new_slide(prs, P())
    add_kicker_title(s, "io.Reader", "The request body is a reader")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "io.Reader — something you can read bytes from.",
        "io.Writer — something you can write bytes to.",
        "Gin’s c.Request.Body is a reader; ShouldBindJSON reads it into your struct.",
        "That is why interfaces matter: one decode works for files, buffers, and HTTP bodies.",
    ], 18)
    add_notes(s, "Canonical story. No need to implement Reader.")

    s = new_slide(prs, P())
    add_kicker_title(s, "In-memory store", "map[int]Student is enough this week", 26)
    add_code(s, MX, 2.05, 12, 2.2, """var (
    mu       sync.Mutex // optional; week 12 locks for real
    students = map[int]models.Student{}
    nextID   = 1
)""", 15)
    add_bullets(s, MX, 4.5, 12, 2.2, [
        "Create assigns nextID, stores, increments.",
        "List ranges the map (order not promised — sort IDs if tests need order).",
        "Get by id: comma-ok. Missing → 404.",
    ], 16)
    add_notes(s, "Mutex optional one-liner. Full race week 12. Slice store also OK.")

    s = new_slide(prs, P())
    add_kicker_title(s, "SMS milestone", "Three routes today")
    add_code(s, MX, 2.05, 12, 2.6, """POST /api/students          → 201 + created student
GET  /api/students          → 200 + array
GET  /api/students/:id      → 200 or 404

No JWT. No Postgres.""", 15)
    add_bullets(s, MX, 4.9, 12, 1.8, [
        "Request body uses first_name, last_name, email, course, …",
        'Bad JSON / missing required fields → 400 with {"error":"..."}.',
    ], 18)
    add_notes(s, "Board the three routes. Homework adds PUT + DELETE.")

    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "models", "Add internal/models/student.go", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 4.0, """package models
type Student struct {
    ID        int    `json:"id"`
    FirstName string `json:"first_name"`
    LastName  string `json:"last_name"`
    Email     string `json:"email"`
    Phone     string `json:"phone"`
    Course    string `json:"course"`
}""", 14)
    add_notes(s, "Skip time.Time DOB in live if slow; string DOB ok for week 7.")

    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "Create", "POST /api/students", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 4.0, """func (h *StudentHandler) Create(c *gin.Context) {
    var req models.Student
    if err := c.ShouldBindJSON(&req); err != nil {
        c.JSON(400, gin.H{"error": "invalid json"})
        return
    }
    // assign id, store in map, c.JSON(201, saved)
}""", 14)
    add_notes(s, "Live Postman POST. Show 400 with broken JSON.")

    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "List + get", "GET collection and by id", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 3.8, """r.POST("/api/students", student.Create)
r.GET("/api/students", student.List)
r.GET("/api/students/:id", student.GetByID)

// GetByID: id, err := strconv.Atoi(c.Param("id"))
// s, ok := students[id]; if !ok { 404 }""", 14)
    add_notes(s, "curl list after two creates. Order may vary — say so.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Postman", "Collection grows")
    add_bullets(s, MX, 2.05, 12, 2.0, [
        "Folder Week 7: Create student, List students, Get by id.",
        "Create: Body → raw → JSON with first_name etc.",
        "Save an example 201 response. Homework will add Update / Delete.",
    ], 16)
    add_code(s, MX, 4.2, 12, 2.4, """{
  "first_name": "Ada",
  "last_name": "Lovelace",
  "email": "ada@example.com",
  "phone": "0770000000",
  "course": "SE"
}""", 14)
    add_notes(s, "Walk the room on JSON typos and trailing commas.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Validation", "Reject junk with 400")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "ShouldBindJSON fails → 400.",
        "Empty first_name / email → 400 with a clear error string.",
        "Do not panic. Do not return 500 for bad client input.",
        'Gin binding:"required" is fine if you introduce a request struct.',
    ], 18)
    add_notes(s, "One live 400 demo is enough.")

    s = new_slide(prs, P())
    add_kicker_title(s, "404", "Missing id is not a server crash")
    add_code(s, MX, 2.05, 12, 2.8, """s, ok := students[id]
if !ok {
    c.JSON(http.StatusNotFound, gin.H{"error": "student not found"})
    return
}
c.JSON(http.StatusOK, s)""", 15)
    add_bullets(s, MX, 5.1, 12, 1.6, [
        "Comma-ok again — same idea as week 3 maps.",
        "Wrong id type (/api/students/abc) → 400.",
    ], 18)
    add_notes(s, "Connect to maps literacy.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Lab · rest of class", "Done when all of this is true")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "models.Student with JSON tags exists",
        "Postman POST creates a student (201)",
        "Postman GET list shows it",
        "Get by id works; unknown id → 404",
        "Broken JSON → 400",
        "Health still green; commit / PR started",
    ], 18)
    add_notes(s, "Gate on board. Update/delete not required to leave.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Homework", "Update + delete")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "PUT /api/students/:id — replace fields; missing → 404; bad JSON → 400.",
        "DELETE /api/students/:id — remove; missing → 404.",
        "Postman screenshots or collection export in the PR.",
        "Still no JWT / Postgres.",
    ], 20)
    add_notes(s, "Homework sheet has the checklist.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Next session", "Week 8 — pointers, memory, repository layer", 26)
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Move the map into a StudentRepository.",
        "Handlers call the repo — HTTP must not own storage forever.",
        "Bring working create + list (and homework update/delete if done).",
    ], 20)
    add_notes(s, "Do not extract repository today unless someone finishes early.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Words to keep", "Say these correctly next week")
    add_card(s, MX, 2.1, 5.9, 1.8, "Struct", "Grouped fields under one type.")
    add_card(s, 7.0, 2.1, 5.7, 1.8, "JSON tag", "Renames fields on the wire.")
    add_card(s, MX, 4.1, 5.9, 1.8, "Interface", "Set of methods; error is one.")
    add_card(s, 7.0, 4.1, 5.7, 1.8, "In-memory", "Store in process RAM — lost on restart.")
    add_notes(s, "Oral quiz.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Layout reminder", "Where the new files live")
    add_code(s, MX, 2.05, 12, 4.2, """backend/
  cmd/server/main.go
  internal/models/student.go
  internal/handlers/
    health.go
    hello.go
    student.go          # Create, List, GetByID (+ homework)
    name.go""", 15)
    add_notes(s, "Store can live in student handler file for now.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Honest limits", "What this week is not")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Not durable — restart wipes students.",
        "Not multi-user safe yet — races possible under load (week 12).",
        "Not authenticated — anyone can POST (week 15).",
        "That is intentional spiral teaching.",
    ], 20)
    add_notes(s, "Set expectations; avoid shame about 'toy' store.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Safety", "Still true with POST bodies")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Never log passwords (none today — good).",
        "Do not commit real emails of classmates without consent — use example.com.",
        "gofmt; no binaries; .env still ignored.",
    ], 20)
    add_notes(s, "Quick.")

    s = new_slide(prs, P(), dark=True)
    add_textbox(s, MX, 3.0, 12, 1.2, "Questions", size=64, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_notes(s, "Then models → create → list → Postman lab.")

    if page != TOTAL:
        raise SystemExit(f"expected {TOTAL} slides, built {page}")
    prs.save(str(dest))
    print("Wrote", dest)
    return dest


if __name__ == "__main__":
    build()

