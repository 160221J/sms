#!/usr/bin/env python3
"""Build an editable PowerPoint of the Week 6 Day 1 deck (native text, not screenshots)."""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    req = Path(__file__).resolve().parent / "requirements.txt"
    sys.stderr.write(
        "Missing python-pptx. From the repo root run:\n"
        f"  python3 -m pip install --user -r {req}\n"
    )
    raise

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[1]
LOGO = REPO / "docs" / "branding" / "epic-learn-logo.png"
OUT = ROOT / "pdf" / "week-06-day-01-slides.pptx"
FOOTER = "Epic Learn Institute of Higher Education - Go Programming Master Course"

INK = RGBColor(0x14, 0x20, 0x2B)
MUTED = RGBColor(0x4D, 0x5D, 0x6B)
TEAL = RGBColor(0x0B, 0x6E, 0x68)
TEAL_DARK = RGBColor(0x08, 0x4F, 0x4B)
NAVY = RGBColor(0x0A, 0x2A, 0x5C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF4, 0xEF, 0xE6)
CARD = RGBColor(0xFF, 0xFD, 0xF9)
LINE = RGBColor(0xD8, 0xCF, 0xC2)
CODE_BG = RGBColor(0x10, 0x21, 0x2C)
CODE_FG = RGBColor(0xE8, 0xF2, 0xEF)
GOLD = RGBColor(0xD9, 0xB8, 0x26)
TITLE_BG = RGBColor(0x14, 0x2A, 0x32)
CREAM_TEXT = RGBColor(0xEA, 0xD9, 0xA3)
LIVE = RGBColor(0xC2, 0x4E, 0x1D)

W, H = 13.333, 7.5
MX = 0.65
HEADER_H = 0.78
FOOTER_H = 0.42
CONTENT_TOP = 0.95
TOTAL = 24


def _font(run, name="Calibri", size=18, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_textbox(slide, l, t, w, h, text, size=18, bold=False, color=INK, font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _font(run, font, size, bold, color)
    return box


def add_bullets(slide, l, t, w, h, items, size=20, color=INK):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        p.level = 0
        run = p.add_run()
        run.text = "•  " + item
        _font(run, "Calibri", size, False, color)
    return box


def add_code(slide, l, t, w, h, text, size=15):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = CODE_BG
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.1)
    lines = text.strip("\n").split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line if line else " "
        _font(run, "Consolas", size, False, CODE_FG)
    return shape


def add_card(slide, l, t, w, h, title, body, title_size=16, body_size=13):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = LINE
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    _font(run, "Calibri", title_size, True, TEAL_DARK)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    run2 = p2.add_run()
    run2.text = body
    _font(run2, "Calibri", body_size, False, INK)
    return shape


def add_brand(slide, page, dark=False):
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(HEADER_H))
    _fill(header, WHITE)
    slide.shapes.add_picture(str(LOGO), Inches(0.35), Inches(0.12), height=Inches(0.55))
    add_textbox(
        slide, 10.2, 0.22, 2.8, 0.4,
        "WEEK 6  ·  DAY 1",
        size=12, bold=True, color=CREAM_TEXT if dark else TEAL_DARK, align=PP_ALIGN.RIGHT,
    )
    gold = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(HEADER_H), Inches(W - 0.7), Inches(0.02)
    )
    _fill(gold, GOLD)

    footer_bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(H - FOOTER_H), Inches(W), Inches(FOOTER_H)
    )
    _fill(footer_bg, WHITE)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(H - FOOTER_H), Inches(W - 0.7), Inches(0.015)
    )
    _fill(line, NAVY)
    add_textbox(
        slide, 0.5, H - FOOTER_H + 0.06, 10.5, 0.3,
        FOOTER, size=11, color=NAVY, align=PP_ALIGN.LEFT,
    )
    add_textbox(
        slide, 11.3, H - FOOTER_H + 0.06, 1.6, 0.3,
        f"{page} / {TOTAL}", size=11, color=MUTED, align=PP_ALIGN.RIGHT,
    )


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_kicker_title(slide, kicker, title, title_size=32, top=None):
    top = CONTENT_TOP if top is None else top
    add_textbox(slide, MX, top, 12, 0.32, kicker.upper(), size=13, bold=True, color=TEAL)
    add_textbox(slide, MX, top + 0.28, 12, 0.7, title, size=title_size, bold=True, color=INK)


def new_slide(prs, page, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if dark:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
        _fill(bg, TITLE_BG)
    else:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
        _fill(bg, CREAM)
    add_brand(slide, page, dark=dark)
    return slide




def build(dest: Path | None = None) -> Path:
    dest = dest or OUT
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    page = 0

    def P():
        nonlocal page
        page += 1
        return page

    s = new_slide(prs, P(), dark=True)
    add_textbox(s, MX, 2.15, 12, 0.4, "GO PROGRAMMING MASTER COURSE", size=14, bold=True, color=CREAM_TEXT)
    add_textbox(s, MX, 2.55, 12, 1.0, "Week 6, Day 1", size=48, bold=True, color=WHITE)
    add_textbox(s, MX, 3.7, 11, 1.1, "Functions and methods inside SMS — extract handlers from main, keep health green.", size=22, color=RGBColor(0xE7, 0xDD, 0xD0))
    add_textbox(s, MX, 5.3, 11, 0.9, "Umesh Indrajith · BSc Eng (Hons), University of Moratuwa\nSenior Software Engineer / Lecturer", size=16, color=RGBColor(0xD7, 0xCB, 0xB8))
    add_notes(s, "Refactor week. No new product feature required.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Since last week", "A running API is the starting line")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "GET /health returns 200 with the exact JSON. If not, fix that first.",
        "Cold-call: what is a handler? What does Gin wrap?",
        "Today we refactor, not add Postgres or JWT.",
        "gofmt still wins. Unformatted PRs returned.",
    ], 20)
    add_notes(s, "Spot-check three health endpoints.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Today", "If Day 1 of week 6 works, you can")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Name the parts of a function: name, params, results.",
        "Return (value, error) and never silently ignore err.",
        "Use defer for cleanup at function exit.",
        "Tell a function from a method; preview pointer receivers.",
        "Move health/hello into internal/handlers; keep :8080 green.",
        "Ship FullName with a test.",
    ], 18)
    add_notes(s, "Deep structs week 7.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Software engineering", "Refactor is a feature when the layout was lying", 26)
    add_textbox(s, MX, 2.1, 12, 1.1, "main that owns every route will rot. Handlers belong in a package with one job: talk HTTP.", size=18, color=MUTED)
    add_bullets(s, MX, 3.4, 12, 3.0, [
        "SMS milestone: extract, do not expand the product surface.",
        "PR description says what moved and why.",
        "Green health before and after is the proof.",
    ], 18)
    add_notes(s, "Connect to week 4 layers.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Functions", "Parts of a function")
    add_code(s, MX, 2.05, 12, 2.2, """func FullName(first, last string) string {
    return first + " " + last
}
func Max(nums []int) (int, bool) { /* name · params · results */ }""", 15)
    add_bullets(s, MX, 4.5, 12, 2.2, [
        "Name — exported if capital (FullName).",
        "Parameters — typed; same type can share.",
        "Results — zero, one, or many. Named results: sparingly.",
    ], 18)
    add_notes(s, "Reuse week 2/3 FullName story.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Pass by value", "Arguments are copies — unless you share memory", 26)
    add_code(s, MX, 2.05, 12, 2.0, """func bump(n int) { n++ }
x := 3
bump(x)
fmt.Println(x) // still 3""", 16)
    add_bullets(s, MX, 4.3, 12, 2.4, [
        "Ints and structs (by default) are copied into the function.",
        "Slices/maps copy the header but share backing data.",
        "Methods that change state usually use a pointer receiver (preview).",
    ], 18)
    add_notes(s, "Plant *StudentHandler need.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Multiple returns", "return x, err is the Go handshake", 28)
    add_code(s, MX, 2.05, 12, 2.6, """port, err := ParsePort(os.Getenv("PORT"))
if err != nil {
    log.Fatal(err)
}""", 16)
    add_bullets(s, MX, 4.9, 12, 1.8, [
        "Ignoring err with _ is a bug unless you can explain why.",
        "Handlers return JSON errors — still check err from callees.",
    ], 18)
    add_notes(s, "Week 9 deepens errors.")

    s = new_slide(prs, P())
    add_kicker_title(s, "defer", "Run this when the function exits")
    add_code(s, MX, 2.05, 12, 2.4, """f, err := os.Open("notes.txt")
if err != nil { return }
defer f.Close() // runs at end of function""", 16)
    add_bullets(s, MX, 4.7, 12, 2.0, [
        "LIFO: last defer runs first.",
        "Use for close, unlock, and later recover patterns.",
    ], 18)
    add_notes(s, "Do not teach recover today.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Methods", "A function with a receiver")
    add_code(s, MX, 2.05, 12, 2.8, """type HealthHandler struct{}
func (h *HealthHandler) Get(c *gin.Context) {
    c.JSON(http.StatusOK, gin.H{"status": "ok", "message": "Backend is running"})
}
h := &HealthHandler{}
r.GET("/health", h.Get)""", 14)
    add_bullets(s, MX, 5.1, 12, 1.6, [
        "Course target: (h *StudentHandler) style in internal/handlers.",
    ], 18)
    add_notes(s, "Struct fields can stay empty today.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Constructors", "New… is just a function that returns a value", 26)
    add_code(s, MX, 2.05, 12, 2.4, """func NewHelloHandler() *HelloHandler {
    return &HelloHandler{}
}""", 16)
    add_bullets(s, MX, 4.7, 12, 2.0, [
        "Copy the shape — full struct theory is week 7.",
        "Later NewStudentHandler(repo) injects dependencies.",
    ], 18)
    add_notes(s, "Toward this repo's handlers.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Pointer vs value · preview", "Why handlers use *HealthHandler", 26)
    add_card(s, MX, 2.1, 5.9, 1.8, "Value receiver", "func (h HealthHandler) — method gets a copy.")
    add_card(s, 7.0, 2.1, 5.7, 1.8, "Pointer receiver", "func (h *HealthHandler) — shared instance; holds repo later.")
    add_card(s, MX, 4.1, 5.9, 1.8, "Course default", "HTTP handlers: pointer receivers.")
    add_card(s, 7.0, 4.1, 5.7, 1.8, "Deep dive", "Stack/heap — week 8. Memorise the pattern today.")
    add_notes(s, "Do not deep-teach pointers.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Generics · one example", "Then stop")
    add_code(s, MX, 2.05, 12, 2.6, """func First[T any](s []T) (T, bool) {
    var zero T
    if len(s) == 0 { return zero, false }
    return s[0], true
}""", 15)
    add_bullets(s, MX, 4.9, 12, 1.8, [
        "One example so the syntax is not alien.",
        "Do not build a generics library this week.",
    ], 18)
    add_notes(s, "Pace valve.")

    s = new_slide(prs, P())
    add_kicker_title(s, "SMS milestone", "Extract week 5 into internal/handlers", 26)
    add_code(s, MX, 2.0, 12, 3.2, """backend/
  cmd/server/main.go       # wire + Run
  internal/handlers/
    health.go
    hello.go
    name.go + name_test.go""", 15)
    add_bullets(s, MX, 5.4, 12, 1.3, [
        "Health JSON contract must not change.",
    ], 18)
    add_notes(s, "Leave on projector during lab.")

    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "Refactor", "Move health out of main", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 3.8, """// handlers: NewHealthHandler + (h *HealthHandler) Get
health := handlers.NewHealthHandler()
r.GET("/health", health.Get)
r.Run(":8080")""", 16)
    add_notes(s, "Live type. curl after.")

    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "Helper", "FullName for later student responses", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 3.6, """func FullName(first, last string) string {
    first = strings.TrimSpace(first)
    last = strings.TrimSpace(last)
    // empty rules + one space
}
// go test ./internal/handlers""", 15)
    add_notes(s, "Red → green once.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Wiring", "What main should look like after")
    add_code(s, MX, 2.05, 12, 3.0, """func main() {
    r := gin.Default()
    r.GET("/health", handlers.NewHealthHandler().Get)
    r.GET("/api/hello", handlers.NewHelloHandler().Get)
    r.Run(":8080")
}""", 15)
    add_bullets(s, MX, 5.3, 12, 1.4, [
        "No business logic in main beyond wiring.",
    ], 18)
    add_notes(s, "Optional routes.Register.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Ignoring errors", "_ is honest; silence is not")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Bad: call a function that returns error and ignore it.",
        "Blank identifier means you chose not to use a value.",
        "Reviewers will ask about ignored error returns.",
    ], 20)
    add_notes(s, "Short slide.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Lab gate", "Done when all of this is true")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Health + hello still 200 on :8080",
        "Handler logic under internal/handlers",
        "FullName + green go test ./internal/handlers",
        "main mostly wires routes",
        "PR opened describing the refactor",
    ], 18)
    add_notes(s, "Checklist on board.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Homework", "PR the refactor — no new product feature")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Required: handlers package + thin main + FullName tests.",
        "README: note the new layout in one paragraph.",
        "Stretch: internal/routes Register helper.",
        "Do not add JWT, Postgres, or student CRUD.",
    ], 20)
    add_notes(s, "If refactor is real, no feature required.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Next session", "Week 7 — structs, JSON tags, in-memory students")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "models.Student, create + list via Postman.",
        "JSON tags matching this repo.",
        "Bring a clean handlers layout and green health.",
    ], 20)
    add_notes(s, "No CRUD today.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Words to keep", "Say these correctly next week")
    add_card(s, MX, 2.1, 5.9, 1.8, "Function", "Named code with params and results.")
    add_card(s, 7.0, 2.1, 5.7, 1.8, "Method", "Function with a receiver type.")
    add_card(s, MX, 4.1, 5.9, 1.8, "defer", "Runs at function exit (LIFO).")
    add_card(s, 7.0, 4.1, 5.7, 1.8, "Constructor", "NewX() returning a ready value.")
    add_notes(s, "Oral quiz if time.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Review bar", "Instructor merges when")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "gofmt clean; health JSON unchanged",
        "Handlers not still pasted entirely inside main",
        "go test ./internal/handlers green",
        "PR says what moved — not a zip dump",
    ], 20)
    add_notes(s, "Collect PR URLs.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Safety", "Still true while refactoring")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "No secrets in git.",
        "Do not “fix” by deleting go.mod.",
        "Keep port :8080 unless the class agrees otherwise.",
        "Commit working health before a big move — small steps.",
    ], 20)
    add_notes(s, "Encourage commit after each extract.")

    s = new_slide(prs, P(), dark=True)
    add_textbox(s, MX, 3.0, 12, 1.2, "Questions", size=64, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_notes(s, "Then live extract → FullName test → lab.")

    if page != TOTAL:
        raise SystemExit(f"expected {TOTAL} slides, built {page}")
    prs.save(str(dest))
    print("Wrote", dest)
    return dest


if __name__ == "__main__":
    build()
