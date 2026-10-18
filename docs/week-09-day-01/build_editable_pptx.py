#!/usr/bin/env python3
"""Build an editable PowerPoint of the Week 9 Day 1 deck (native text, not screenshots)."""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    req = Path(__file__).resolve().parent / "requirements.txt"
    sys.stderr.write(
        "Missing python-pptx. From the repo root run:\n"
        f"  python3 -m pip install --user -r {req}\n"
    )
    raise

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[1]
LOGO = REPO / "docs" / "branding" / "epic-learn-logo.png"
OUT = ROOT / "pdf" / "week-09-day-01-slides.pptx"
FOOTER = "Epic Learn Institute of Higher Education - Go Programming Master Course"

INK = RGBColor(0x14, 0x20, 0x2B)
MUTED = RGBColor(0x4D, 0x5D, 0x6B)
TEAL = RGBColor(0x0B, 0x6E, 0x68)
TEAL_DARK = RGBColor(0x08, 0x4F, 0x4B)
NAVY = RGBColor(0x0A, 0x2A, 0x5C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF4, 0xEF, 0xE6)
CARD = RGBColor(0xFF, 0xFD, 0xF9)
LINE = RGBColor(0xD8, 0xCF, 0xC2)
CODE_BG = RGBColor(0x10, 0x21, 0x2C)
CODE_FG = RGBColor(0xE8, 0xF2, 0xEF)
GOLD = RGBColor(0xD9, 0xB8, 0x26)
TITLE_BG = RGBColor(0x14, 0x2A, 0x32)
CREAM_TEXT = RGBColor(0xEA, 0xD9, 0xA3)
LIVE = RGBColor(0xC2, 0x4E, 0x1D)

W, H = 13.333, 7.5
MX = 0.65
HEADER_H = 0.78
FOOTER_H = 0.42
CONTENT_TOP = 0.95
TOTAL = 26


def _font(run, name="Calibri", size=18, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_textbox(slide, l, t, w, h, text, size=18, bold=False, color=INK, font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _font(run, font, size, bold, color)
    return box


def add_bullets(slide, l, t, w, h, items, size=20, color=INK):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        p.level = 0
        run = p.add_run()
        run.text = "•  " + item
        _font(run, "Calibri", size, False, color)
    return box


def add_code(slide, l, t, w, h, text, size=15):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = CODE_BG
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.1)
    lines = text.strip("\n").split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line if line else " "
        _font(run, "Consolas", size, False, CODE_FG)
    return shape


def add_card(slide, l, t, w, h, title, body, title_size=16, body_size=13):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = LINE
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    _font(run, "Calibri", title_size, True, TEAL_DARK)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    run2 = p2.add_run()
    run2.text = body
    _font(run2, "Calibri", body_size, False, INK)
    return shape


def add_brand(slide, page, dark=False):
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(HEADER_H))
    _fill(header, WHITE)
    slide.shapes.add_picture(str(LOGO), Inches(0.35), Inches(0.12), height=Inches(0.55))
    add_textbox(
        slide, 10.2, 0.22, 2.8, 0.4,
        "WEEK 9  ·  DAY 1",
        size=12, bold=True, color=CREAM_TEXT if dark else TEAL_DARK, align=PP_ALIGN.RIGHT,
    )
    gold = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(HEADER_H), Inches(W - 0.7), Inches(0.02)
    )
    _fill(gold, GOLD)

    footer_bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(H - FOOTER_H), Inches(W), Inches(FOOTER_H)
    )
    _fill(footer_bg, WHITE)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(H - FOOTER_H), Inches(W - 0.7), Inches(0.015)
    )
    _fill(line, NAVY)
    add_textbox(
        slide, 0.5, H - FOOTER_H + 0.06, 10.5, 0.3,
        FOOTER, size=11, color=NAVY, align=PP_ALIGN.LEFT,
    )
    add_textbox(
        slide, 11.3, H - FOOTER_H + 0.06, 1.6, 0.3,
        f"{page} / {TOTAL}", size=11, color=MUTED, align=PP_ALIGN.RIGHT,
    )


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_kicker_title(slide, kicker, title, title_size=32, top=None):
    top = CONTENT_TOP if top is None else top
    add_textbox(slide, MX, top, 12, 0.32, kicker.upper(), size=13, bold=True, color=TEAL)
    add_textbox(slide, MX, top + 0.28, 12, 0.7, title, size=title_size, bold=True, color=INK)


def new_slide(prs, page, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if dark:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
        _fill(bg, TITLE_BG)
    else:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
        _fill(bg, CREAM)
    add_brand(slide, page, dark=dark)
    return slide


def build(dest: Path | None = None) -> Path:
    dest = dest or OUT
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    page = 0

    def P():
        nonlocal page
        page += 1
        return page

    s = new_slide(prs, P(), dark=True)
    add_textbox(s, MX, 2.15, 12, 0.4, "GO PROGRAMMING MASTER COURSE", size=14, bold=True, color=CREAM_TEXT)
    add_textbox(s, MX, 2.55, 12, 1.0, "Week 9, Day 1", size=48, bold=True, color=WHITE)
    add_textbox(s, MX, 3.7, 11, 1.1, "Errors, wrapping, validation, slog — and a light register/login preview.", size=22, color=RGBColor(0xE7, 0xDD, 0xD0))
    add_textbox(s, MX, 5.3, 11, 0.9, "Umesh Indrajith · BSc Eng (Hons), University of Moratuwa\nSenior Software Engineer / Lecturer", size=16, color=RGBColor(0xD7, 0xCB, 0xB8))
    add_notes(s, "Week 8 handler→repo must exist. Today: ErrNotFound→404, slog, config.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Since last week", "Repo owns the map. Handlers own HTTP status.", 26)
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Create / list / get still green through StudentRepository.",
        "Cold-call: what does ok == false mean in GetByID? Why not panic?",
        "Today we replace bool with real error values — and log properly.",
        "JWT middleware can wait; prefer register/login without locking students.",
    ], 18)
    add_notes(s, "Fix broken week-8 extracts in first 10 minutes only.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Today", "If Day 1 of week 9 works, you can")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Treat error as a value; use a sentinel ErrNotFound.",
        "Wrap with %w; check with errors.Is / errors.As.",
        'Never panic in an HTTP handler — return {"error":"..."}.',
        "Validate with Gin binding; log 5xx with log/slog.",
        "Load config from env (.env.example, never commit secrets).",
    ], 18)
    add_notes(s, "Lab: ErrNotFound map. Homework: slog + .env.example.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Software engineering", "Errors are part of the API contract", 26)
    add_textbox(s, MX, 2.1, 12, 1.0, "Clients and operators need predictable JSON and logs. A panic is a bug, not a status code.", size=18, color=MUTED)
    add_bullets(s, MX, 3.3, 12, 3.2, [
        'Same shape everywhere: {"error":"student not found"}.',
        "4xx = client problem. 5xx = our problem — and we log it.",
        "Secrets live in environment variables, not in git.",
    ], 18)
    add_notes(s, "Write the JSON contract on the board.")

    s = new_slide(prs, P())
    add_kicker_title(s, "error as a value", "The second return is the handshake")
    add_code(s, MX, 2.05, 12, 2.6, """s, err := repo.GetByID(id)
if err != nil {
    // decide: 404, 400, or 500 — do not panic
    return
}
c.JSON(http.StatusOK, s)""", 15)
    add_bullets(s, MX, 4.9, 12, 1.8, [
        "Ignoring err with _ needs a spoken reason.",
        "error is an interface: anything with Error() string.",
    ], 18)
    add_notes(s, "Connect to week 6/7 multiple returns.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Sentinel vs custom", "Name the expected failures")
    add_code(s, MX, 2.05, 12, 2.4, """var ErrNotFound = errors.New("not found")

type ValidationError struct { Field, Msg string }
func (e ValidationError) Error() string { return e.Field + ": " + e.Msg }""", 14)
    add_bullets(s, MX, 4.7, 12, 2.0, [
        "Sentinel: compare with errors.Is.",
        "Custom type: when you need fields; unwrap with errors.As.",
        "Lab uses a sentinel ErrNotFound.",
    ], 16)
    add_notes(s, "Prefer sentinel for not-found this week.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Wrapping", "%w, Is, and As")
    add_code(s, MX, 2.05, 12, 2.8, """if err != nil {
    return fmt.Errorf("get student %d: %w", id, err)
}
if errors.Is(err, apperrors.ErrNotFound) {
    c.JSON(404, gin.H{"error": "student not found"})
    return
}""", 14)
    add_bullets(s, MX, 5.1, 12, 1.6, [
        "%w preserves the chain so callers can still Is the root cause.",
    ], 18)
    add_notes(s, "Live one wrap + Is.")

    s = new_slide(prs, P())
    add_kicker_title(s, "HTTP mapping", "Repo returns errors · handler picks status", 26)
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "errors.Is(err, ErrNotFound) → 404",
        "Bad input / bind failure → 400",
        "Anything unexpected → 500 + slog.Error",
        "Never panic(err) inside a Gin handler.",
    ], 20)
    add_notes(s, "Board this table.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Validation", "Reject junk before it hits the repo")
    add_code(s, MX, 2.05, 12, 2.8, """type CreateStudentRequest struct {
    FirstName string `json:"first_name" binding:"required"`
    Email     string `json:"email" binding:"required,email"`
    Course    string `json:"course" binding:"required"`
}""", 14)
    add_bullets(s, MX, 5.1, 12, 1.6, [
        "ShouldBindJSON failure → 400 with {\"error\":...}.",
        "Match this repo’s RegisterRequest style.",
    ], 18)
    add_notes(s, "Show one 400 with missing email.")

    s = new_slide(prs, P())
    add_kicker_title(s, "slog", "Structured logs, not fmt.Println")
    add_code(s, MX, 2.05, 12, 1.8, """slog.Info("student created", "student_id", s.ID)
slog.Error("create failed", "err", err)""", 16)
    add_bullets(s, MX, 4.1, 12, 2.6, [
        "JSON handler: level, msg, attrs.",
        "Log 5xx with context — never log passwords.",
        "Homework: kill leftover print debugging.",
    ], 18)
    add_notes(s, "Set default JSON logger once in main.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Config", "Environment, not hard-coded secrets")
    add_code(s, MX, 2.05, 12, 2.4, """type Config struct {
    Port      string
    JWTSecret string
}
// godotenv.Load() optional; ship .env.example — never commit .env""", 15)
    add_bullets(s, MX, 4.7, 12, 2.0, [
        "Default port :8080 if empty.",
        "DB vars can appear; we still do not connect Postgres this week.",
    ], 18)
    add_notes(s, "Point at backend/internal/config/config.go.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Auth · light", "Hash passwords. Sign a string later.")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "bcrypt: store a hash, never the password.",
        "JWT: a signed claim string — middleware can wait until week 15.",
        "Prefer POST /api/register (+ login) without locking student routes yet.",
    ], 20)
    add_notes(s, "Pace valve.")

    s = new_slide(prs, P())
    add_kicker_title(s, "JWT preview", "One slide — then stop")
    add_code(s, MX, 2.05, 12, 2.0, """// idea only: header.payload.signature
// payload might carry user_id + exp
// verify signature with JWT_SECRET from config""", 16)
    add_bullets(s, MX, 4.3, 12, 2.4, [
        "Product story: login returns a token; later middleware checks it.",
        "Do not protect /api/students today unless the room is flying.",
    ], 18)
    add_notes(s, "Show LoginResponse shape if useful.")

    s = new_slide(prs, P())
    add_kicker_title(s, "SMS milestone", "Consistent errors + logs")
    add_code(s, MX, 2.05, 12, 2.4, """GET /api/students/:id missing → 404 {"error":"student not found"}
Unexpected repo failure → 500 + slog.Error
Optional: POST /api/register (in-memory users)""", 15)
    add_bullets(s, MX, 4.7, 12, 2.0, [
        "Files: auth_handler.go, models/user.go, config/config.go.",
        "Still no Postgres. Restart still clears memory.",
    ], 18)
    add_notes(s, "Board lab vs stretch.")

    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "ErrNotFound", "Sentinel from repo → 404 in handler", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 4.0, """if !ok { return models.Student{}, apperrors.ErrNotFound }

if errors.Is(err, apperrors.ErrNotFound) {
    c.JSON(404, gin.H{"error": "student not found"})
    return
}
slog.Error("get student", "id", id, "err", err)
c.JSON(500, gin.H{"error": "internal error"})""", 14)
    add_notes(s, "Live curl missing id.")

    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "slog + config", "JSON logs and .env.example", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 3.6, """slog.SetDefault(slog.New(slog.NewJSONHandler(os.Stdout, nil)))
cfg := config.Load()
r.Run(":" + cfg.Port)

# .env.example
PORT=8080
JWT_SECRET=change-me""", 15)
    add_notes(s, "Show one JSON log line.")

    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO · STRETCH", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "Register", "POST /api/register in memory", top=CONTENT_TOP + 0.28)
    add_bullets(s, MX, 2.4, 12, 4.0, [
        "bcrypt hash → store User (password_hash omitted in JSON).",
        "Duplicate username → 409.",
        "Do not protect student routes yet.",
        "Skip live if ErrNotFound lab needs the rest of class.",
    ], 18)
    add_notes(s, "Optional.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Lab · rest of class", "Done when all of this is true")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "ErrNotFound defined; repo returns it for missing students",
        "Handler maps not-found → 404; unexpected → 500 + slog",
        "Bind failures → 400 with {\"error\":...}",
        "No panics on the student paths",
        "Health still green; PR started",
    ], 18)
    add_notes(s, "Gate on board.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Homework", "slog + .env.example")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Replace fmt.Println debugging with slog (Info/Error + attrs).",
        "Add .env.example (PORT, JWT_SECRET). Keep .env gitignored.",
        "Wire internal/config for port (and JWT secret for later).",
        "Stretch: in-memory register (+ login token if ready).",
    ], 18)
    add_notes(s, "Homework sheet has the checklist.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Next session", "Week 10 — goroutines and the scheduler", 26)
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Concurrency vs parallelism; go f().",
        "Why handlers must not fire-and-forget without a plan.",
        "Bring clean error mapping and slog.",
    ], 20)
    add_notes(s, "Do not teach channels today.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Words to keep", "Say these correctly next week")
    add_card(s, MX, 2.1, 5.9, 1.8, "Sentinel", "Named error value; compare with errors.Is.")
    add_card(s, 7.0, 2.1, 5.7, 1.8, "Wrap", "fmt.Errorf with %w keeps the cause.")
    add_card(s, MX, 4.1, 5.9, 1.8, "slog", "Structured logs with level + attrs.")
    add_card(s, 7.0, 4.1, 5.7, 1.8, "Config", "Secrets from env; example file only in git.")
    add_notes(s, "Oral quiz.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Layout reminder", "Where the new files live")
    add_code(s, MX, 2.05, 12, 4.0, """backend/
  .env.example
  cmd/server/main.go
  internal/config/config.go
  internal/apperrors/errors.go
  internal/models/{student,user}.go
  internal/handlers/{student,auth}.go""", 14)
    add_notes(s, "Auth handler optional for lab gate.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Honest limits", "What this week is not")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Not Postgres — users/students still in memory.",
        "Not full auth middleware on every route.",
        "Not a logging platform — stdout JSON is enough.",
        "Not panicking “to be honest” in handlers.",
    ], 18)
    add_notes(s, "Keep spiral honest.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Safety", "Secrets and passwords")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Never commit .env or real JWT secrets.",
        "Never log raw passwords or full tokens.",
        "gofmt; no binaries in git.",
    ], 20)
    add_notes(s, "Quick.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Review bar", "Instructor merges when")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Missing student → 404 JSON, not panic / empty 200",
        "Unexpected errors → 500 + slog (evidence in PR)",
        ".env.example present; .env absent from git",
        "gofmt clean; PR describes error mapping",
    ], 18)
    add_notes(s, "Collect PR URLs.")

    s = new_slide(prs, P(), dark=True)
    add_textbox(s, MX, 3.0, 12, 1.2, "Questions", size=64, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_notes(s, "Then ErrNotFound live → slog → lab.")

    if page != TOTAL:
        raise SystemExit(f"expected {TOTAL} slides, built {page}")
    prs.save(str(dest))
    print("Wrote", dest)
    return dest


if __name__ == "__main__":
    build()
