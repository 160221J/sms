#!/usr/bin/env python3
"""Build an editable PowerPoint of the Week 4 Day 1 deck (native text, not screenshots)."""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    req = Path(__file__).resolve().parent / "requirements.txt"
    sys.stderr.write(
        "Missing python-pptx. From the repo root run:\n"
        f"  python3 -m pip install --user -r {req}\n"
    )
    raise

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[1]
LOGO = REPO / "docs" / "branding" / "epic-learn-logo.png"
OUT = ROOT / "pdf" / "week-04-day-01-slides.pptx"
FOOTER = "Epic Learn Institute of Higher Education - Go Programming Master Course"

INK = RGBColor(0x14, 0x20, 0x2B)
MUTED = RGBColor(0x4D, 0x5D, 0x6B)
TEAL = RGBColor(0x0B, 0x6E, 0x68)
TEAL_DARK = RGBColor(0x08, 0x4F, 0x4B)
NAVY = RGBColor(0x0A, 0x2A, 0x5C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF4, 0xEF, 0xE6)
CARD = RGBColor(0xFF, 0xFD, 0xF9)
LINE = RGBColor(0xD8, 0xCF, 0xC2)
CODE_BG = RGBColor(0x10, 0x21, 0x2C)
CODE_FG = RGBColor(0xE8, 0xF2, 0xEF)
GOLD = RGBColor(0xD9, 0xB8, 0x26)
TITLE_BG = RGBColor(0x14, 0x2A, 0x32)
CREAM_TEXT = RGBColor(0xEA, 0xD9, 0xA3)
LIVE = RGBColor(0xC2, 0x4E, 0x1D)

W, H = 13.333, 7.5
MX = 0.65
HEADER_H = 0.78
FOOTER_H = 0.42
CONTENT_TOP = 0.95
TOTAL = 24


def _font(run, name="Calibri", size=18, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_textbox(slide, l, t, w, h, text, size=18, bold=False, color=INK, font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _font(run, font, size, bold, color)
    return box


def add_bullets(slide, l, t, w, h, items, size=20, color=INK):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        p.level = 0
        run = p.add_run()
        run.text = "•  " + item
        _font(run, "Calibri", size, False, color)
    return box


def add_code(slide, l, t, w, h, text, size=15):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = CODE_BG
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.1)
    lines = text.strip("\n").split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line if line else " "
        _font(run, "Consolas", size, False, CODE_FG)
    return shape


def add_card(slide, l, t, w, h, title, body, title_size=16, body_size=13):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = LINE
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    _font(run, "Calibri", title_size, True, TEAL_DARK)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    run2 = p2.add_run()
    run2.text = body
    _font(run2, "Calibri", body_size, False, INK)
    return shape


def add_brand(slide, page, dark=False):
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(HEADER_H))
    _fill(header, WHITE)
    slide.shapes.add_picture(str(LOGO), Inches(0.35), Inches(0.12), height=Inches(0.55))
    add_textbox(
        slide, 10.2, 0.22, 2.8, 0.4,
        "WEEK 4  ·  DAY 1",
        size=12, bold=True, color=CREAM_TEXT if dark else TEAL_DARK, align=PP_ALIGN.RIGHT,
    )
    gold = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(HEADER_H), Inches(W - 0.7), Inches(0.02)
    )
    _fill(gold, GOLD)

    footer_bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(H - FOOTER_H), Inches(W), Inches(FOOTER_H)
    )
    _fill(footer_bg, WHITE)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(H - FOOTER_H), Inches(W - 0.7), Inches(0.015)
    )
    _fill(line, NAVY)
    add_textbox(
        slide, 0.5, H - FOOTER_H + 0.06, 10.5, 0.3,
        FOOTER, size=11, color=NAVY, align=PP_ALIGN.LEFT,
    )
    add_textbox(
        slide, 11.3, H - FOOTER_H + 0.06, 1.6, 0.3,
        f"{page} / {TOTAL}", size=11, color=MUTED, align=PP_ALIGN.RIGHT,
    )


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_kicker_title(slide, kicker, title, title_size=32, top=None):
    top = CONTENT_TOP if top is None else top
    add_textbox(slide, MX, top, 12, 0.32, kicker.upper(), size=13, bold=True, color=TEAL)
    add_textbox(slide, MX, top + 0.28, 12, 0.7, title, size=title_size, bold=True, color=INK)


def new_slide(prs, page, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if dark:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
        _fill(bg, TITLE_BG)
    else:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
        _fill(bg, CREAM)
    add_brand(slide, page, dark=dark)
    return slide






def build(dest: Path | None = None) -> Path:
    dest = dest or OUT
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    page = 0

    def P():
        nonlocal page
        page += 1
        return page

    s = new_slide(prs, P(), dark=True)
    add_textbox(s, MX, 2.15, 12, 0.4, "GO PROGRAMMING MASTER COURSE", size=14, bold=True, color=CREAM_TEXT)
    add_textbox(s, MX, 2.55, 12, 1.0, "Week 4, Day 1", size=48, bold=True, color=WHITE)
    add_textbox(s, MX, 3.7, 11, 1.1, "Modules, packages, and project shape — today the Student Management System clock starts.", size=22, color=RGBColor(0xE7, 0xDD, 0xD0))
    add_textbox(s, MX, 5.3, 11, 0.9, "Umesh Indrajith · BSc Eng (Hons), University of Moratuwa\nSenior Software Engineer / Lecturer", size=16, color=RGBColor(0xD7, 0xCB, 0xB8))
    add_notes(s, "Product week 1 of ~12. No HTTP today.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Since last week", "Prove the katas, then start the product")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Week 3 homework: green go test on Contains, Unique, Freq, and friends.",
        "Cold-call: why s = append(s, x)? What is the backing-array gotcha?",
        "gofmt still required. Secrets still never go in git.",
        "Today we start the SMS repo layout.",
    ], 20)
    add_notes(s, "10 minutes max on homework cleanup.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Today", "If Day 1 of week 4 works, you can")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Explain package vs module in one sentence each.",
        "Run go mod init and read a go.mod.",
        "Use internal/ as a boundary, not a random folder name.",
        "Scaffold the SMS backend tree used in this course.",
        "Ship a README and open a clean PR — no .env, no binaries.",
    ], 18)
    add_notes(s, "GET /health is week 5.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Software engineering", "The product clock starts today", 28)
    add_textbox(s, MX, 2.1, 12, 1.2, "Weeks 1–3 were language fuel. From now on every week lands something in SMS.", size=18, color=MUTED)
    add_bullets(s, MX, 3.4, 12, 3.0, [
        "Today: empty module + folders + README. Not an HTTP server yet.",
        "Spiral teaching: just enough layout to grow handlers next week.",
        "One repo per student. Homework is a PR.",
    ], 18)
    add_notes(s, "Draw the 12-week product arc.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Packages", "A package is a folder of Go files with the same name", 26)
    add_code(s, MX, 2.05, 12, 2.2, """package greeter

func Hello(name string) string {
    return "Hello, " + name
}""", 16)
    add_bullets(s, MX, 4.5, 12, 2.2, [
        "All .go files in one directory share one package name.",
        "package main + func main = a program you can go run.",
        "Other packages are libraries you import.",
    ], 18)
    add_notes(s, "Connect to week 3 demo layout.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Imports", "Import path = module path + folder")
    add_code(s, MX, 2.05, 12, 1.8, """import (
    "fmt"
    "student-management-system/internal/models"
)""", 16)
    add_bullets(s, MX, 4.1, 12, 2.6, [
        "Stdlib has short paths: fmt, strings, net/http.",
        "Your code uses the module name from go.mod, then the directory.",
        "Wrong import + no go.mod → package is not in std.",
    ], 18)
    add_notes(s, "Remind classroom gotest/greeter error.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Export rules", "Capital letter = visible outside the package")
    add_card(s, MX, 2.1, 5.9, 1.8, "Exported", "Hello, Student, NewHandler — other packages can use these.")
    add_card(s, 7.0, 2.1, 5.7, 1.8, "Unexported", "hello, parseEnv — private to this package.")
    add_card(s, MX, 4.1, 5.9, 1.8, "Same package", "Files in one folder see each other’s unexported names.")
    add_card(s, 7.0, 4.1, 5.7, 1.8, "Not Java public", "There is no public keyword. The first letter is the rule.")
    add_notes(s, "Quick week-1 recap.")

    s = new_slide(prs, P())
    add_kicker_title(s, "internal/", "Go’s built-in do-not-import-from-outside", 26)
    add_code(s, MX, 2.05, 12, 2.0, """backend/
  cmd/server/          # may import internal/...
  internal/handlers/   # may import internal/models""", 15)
    add_bullets(s, MX, 4.3, 12, 2.4, [
        "internal is a compiler rule, not a style guide.",
        "Put product code in internal/. Put the process entry in cmd/.",
    ], 18)
    add_notes(s, "One sentence: internal stops accidental public APIs.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Modules", "go mod init names the project")
    add_code(s, MX, 2.05, 12, 1.8, """mkdir -p ~/epic-go/sms/backend && cd ~/epic-go/sms/backend
go mod init student-management-system
cat go.mod""", 15)
    add_bullets(s, MX, 4.1, 12, 2.6, [
        "Creates go.mod with a module path — prefixes your imports.",
        "Weeks 2–3 used modules as a ritual. Today we teach what it means.",
        "One module for the SMS backend is enough for this course.",
    ], 18)
    add_notes(s, "Live type this.")

    s = new_slide(prs, P())
    add_kicker_title(s, "go.mod", "The file that defines your module")
    add_code(s, MX, 2.05, 12, 1.6, """module student-management-system

go 1.25""", 18)
    add_bullets(s, MX, 4.0, 12, 2.8, [
        "module line = import prefix for your packages.",
        "go line = language version (pin 1.25.x in class).",
        "require lines appear later when you go get Gin/pgx. Not today.",
    ], 18)
    add_notes(s, "Mention go.sum lightly.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Tools · light", "go get and go mod tidy")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "go get example.com/mod@v1.2.3 — add or upgrade a dependency.",
        "go mod tidy — add missing requires, remove unused ones.",
        "Today you may have zero third-party deps. That is fine.",
        "go mod tidy && go build ./...",
    ], 20)
    add_notes(s, "Pace valve — skip live if late.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Layout", "The SMS backend tree (memorise the jobs)", 26)
    add_code(s, MX, 2.0, 12, 4.6, """backend/
  cmd/server/main.go      # process entry
  internal/config/        # env
  internal/handlers/      # HTTP (week 5+)
  internal/repository/    # data
  internal/models/        # structs
  internal/routes/
  internal/middleware/
  internal/database/
  internal/utils/         # JWT later — tradeoff
  go.mod
  README.md""", 14)
    add_notes(s, "Leave on projector during lab.")

    s = new_slide(prs, P())
    add_kicker_title(s, "cmd/server", "One process entry — keep it thin")
    add_code(s, MX, 2.05, 12, 2.4, """package main
import "fmt"
func main() {
    fmt.Println("sms starting")
}""", 16)
    add_bullets(s, MX, 4.7, 12, 2.0, [
        "cmd/<name> is the usual Go pattern for executables.",
        "Run: go run ./cmd/server from backend/.",
    ], 18)
    add_notes(s, "Lab gate output.")

    s = new_slide(prs, P())
    add_kicker_title(s, "internal packages", "One job per folder")
    add_card(s, MX, 2.1, 5.9, 1.8, "handlers", "Talk HTTP. Request → call something → JSON.")
    add_card(s, 7.0, 2.1, 5.7, 1.8, "repository", "Talk data. No Gin imports here.")
    add_card(s, MX, 4.1, 5.9, 1.8, "models", "Shapes: Student, User. No workflow.")
    add_card(s, 7.0, 4.1, 5.7, 1.8, "config", "Read env. Fail fast later if secrets missing.")
    add_notes(s, "Empty folders OK with .gitkeep.")

    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "Scaffold", "Type the skeleton with me", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 4.0, """cd ~/epic-go/sms/backend
go mod init student-management-system
mkdir -p cmd/server internal/{config,handlers,repository,models,routes,middleware,database,utils}
# write cmd/server/main.go → fmt.Println("sms starting")
gofmt -w cmd/server/main.go
go run ./cmd/server""", 14)
    add_notes(s, "Students type along.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Clean code", "Small files, clear names, no junk drawer")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Package name is a noun for its job: handlers, not stuff.",
        "Prefer many small files over one 800-line main.go.",
        "internal/utils exists for JWT later — tradeoff, not a dumping ground this week.",
        "gofmt on save. Unformatted PRs are returned.",
    ], 20)
    add_notes(s, "Call out utils honestly.")

    s = new_slide(prs, P())
    add_kicker_title(s, "README", "The contract of the repo")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Three sentences: what SMS is, who uses it, what done looks like.",
        "How to run (even if today it only prints sms starting).",
        "Module path, Go version, GitHub link.",
        "A README that says only “final year project” fails.",
    ], 20)
    add_notes(s, "Rewrite a bad README live.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Git / PR", "Homework is a pull request, not a zip")
    add_code(s, MX, 2.05, 12, 2.2, """git checkout -b week-04-skeleton
git add . && git commit -m "Add SMS backend skeleton and README"
git push -u origin week-04-skeleton""", 14)
    add_bullets(s, MX, 4.5, 12, 2.2, [
        "PR description: what and why (three lines).",
        "Instructor merges only if gofmt is clean and .env is absent.",
    ], 18)
    add_notes(s, "Fix missing GitHub before they leave.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Secrets", ".env stays local. .env.example is committed.", 26)
    add_code(s, MX, 2.05, 12, 2.4, """# .gitignore
.env
*.exe

# .env.example
PORT=8080
DATABASE_URL=
JWT_SECRET=""", 15)
    add_bullets(s, MX, 4.7, 12, 2.0, [
        "Empty values in the example. Real secrets never in git.",
        "Create the ignore rule now so week 5 is safe.",
    ], 18)
    add_notes(s, "Merge blocker.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Lab gate", "Done when all of this is true")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "backend/go.mod with module student-management-system",
        "Folder tree matches the slide",
        "go run ./cmd/server prints sms starting",
        "README has three product sentences + how to run",
        ".env gitignored; PR URL shown to instructor",
    ], 18)
    add_notes(s, "Checklist on the board.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Homework", "PR the skeleton before next class")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Same gate as the lab — finished and pushed.",
        "PR title example: Week 4 — SMS backend skeleton.",
        "Optional: why internal/ exists, in your own words.",
        "Do not add Gin or HTTP yet. That is week 5.",
    ], 20)
    add_notes(s, "Deliverable is the PR.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Next session", "Week 5 — HTTP, JSON, Gin, first SMS API")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Stdlib net/http hello, then Gin.",
        "GET /health → JSON {\"status\":\"ok\"}.",
        "Bring a reviewable skeleton PR.",
        "Postman arrives. The API becomes real.",
    ], 20)
    add_notes(s, "Do not start Gin today.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Words to keep", "Say these correctly next week")
    add_card(s, MX, 2.1, 5.9, 1.8, "Package", "Folder of Go files with one package clause.")
    add_card(s, 7.0, 2.1, 5.7, 1.8, "Module", "Versioned unit defined by go.mod; import prefix.")
    add_card(s, MX, 4.1, 5.9, 1.8, "internal/", "Compiler-enforced privacy boundary for app code.")
    add_card(s, 7.0, 4.1, 5.7, 1.8, "cmd/", "Place for main packages / executables.")
    add_notes(s, "Oral quiz if time.")

    s = new_slide(prs, P(), dark=True)
    add_textbox(s, MX, 3.0, 12, 1.2, "Questions", size=64, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_notes(s, "Start scaffold. Collect PR URLs.")

    if page != TOTAL:
        raise SystemExit(f"expected {TOTAL} slides, built {page}")
    prs.save(str(dest))
    print("Wrote", dest)
    return dest


if __name__ == "__main__":
    build()
