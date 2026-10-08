#!/usr/bin/env python3
"""Build an editable PowerPoint of the Week 8 Day 1 deck (native text, not screenshots)."""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    req = Path(__file__).resolve().parent / "requirements.txt"
    sys.stderr.write(
        "Missing python-pptx. From the repo root run:\n"
        f"  python3 -m pip install --user -r {req}\n"
    )
    raise

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[1]
LOGO = REPO / "docs" / "branding" / "epic-learn-logo.png"
OUT = ROOT / "pdf" / "week-08-day-01-slides.pptx"
FOOTER = "Epic Learn Institute of Higher Education - Go Programming Master Course"

INK = RGBColor(0x14, 0x20, 0x2B)
MUTED = RGBColor(0x4D, 0x5D, 0x6B)
TEAL = RGBColor(0x0B, 0x6E, 0x68)
TEAL_DARK = RGBColor(0x08, 0x4F, 0x4B)
NAVY = RGBColor(0x0A, 0x2A, 0x5C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF4, 0xEF, 0xE6)
CARD = RGBColor(0xFF, 0xFD, 0xF9)
LINE = RGBColor(0xD8, 0xCF, 0xC2)
CODE_BG = RGBColor(0x10, 0x21, 0x2C)
CODE_FG = RGBColor(0xE8, 0xF2, 0xEF)
GOLD = RGBColor(0xD9, 0xB8, 0x26)
TITLE_BG = RGBColor(0x14, 0x2A, 0x32)
CREAM_TEXT = RGBColor(0xEA, 0xD9, 0xA3)
LIVE = RGBColor(0xC2, 0x4E, 0x1D)

W, H = 13.333, 7.5
MX = 0.65
HEADER_H = 0.78
FOOTER_H = 0.42
CONTENT_TOP = 0.95
TOTAL = 26


def _font(run, name="Calibri", size=18, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_textbox(slide, l, t, w, h, text, size=18, bold=False, color=INK, font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _font(run, font, size, bold, color)
    return box


def add_bullets(slide, l, t, w, h, items, size=20, color=INK):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        p.level = 0
        run = p.add_run()
        run.text = "•  " + item
        _font(run, "Calibri", size, False, color)
    return box


def add_code(slide, l, t, w, h, text, size=15):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = CODE_BG
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.1)
    lines = text.strip("\n").split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line if line else " "
        _font(run, "Consolas", size, False, CODE_FG)
    return shape


def add_card(slide, l, t, w, h, title, body, title_size=16, body_size=13):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = LINE
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    _font(run, "Calibri", title_size, True, TEAL_DARK)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    run2 = p2.add_run()
    run2.text = body
    _font(run2, "Calibri", body_size, False, INK)
    return shape


def add_brand(slide, page, dark=False):
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(HEADER_H))
    _fill(header, WHITE)
    slide.shapes.add_picture(str(LOGO), Inches(0.35), Inches(0.12), height=Inches(0.55))
    add_textbox(
        slide, 10.2, 0.22, 2.8, 0.4,
        "WEEK 8  ·  DAY 1",
        size=12, bold=True, color=CREAM_TEXT if dark else TEAL_DARK, align=PP_ALIGN.RIGHT,
    )
    gold = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(HEADER_H), Inches(W - 0.7), Inches(0.02)
    )
    _fill(gold, GOLD)

    footer_bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(H - FOOTER_H), Inches(W), Inches(FOOTER_H)
    )
    _fill(footer_bg, WHITE)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(H - FOOTER_H), Inches(W - 0.7), Inches(0.015)
    )
    _fill(line, NAVY)
    add_textbox(
        slide, 0.5, H - FOOTER_H + 0.06, 10.5, 0.3,
        FOOTER, size=11, color=NAVY, align=PP_ALIGN.LEFT,
    )
    add_textbox(
        slide, 11.3, H - FOOTER_H + 0.06, 1.6, 0.3,
        f"{page} / {TOTAL}", size=11, color=MUTED, align=PP_ALIGN.RIGHT,
    )


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_kicker_title(slide, kicker, title, title_size=32, top=None):
    top = CONTENT_TOP if top is None else top
    add_textbox(slide, MX, top, 12, 0.32, kicker.upper(), size=13, bold=True, color=TEAL)
    add_textbox(slide, MX, top + 0.28, 12, 0.7, title, size=title_size, bold=True, color=INK)


def new_slide(prs, page, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if dark:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
        _fill(bg, TITLE_BG)
    else:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
        _fill(bg, CREAM)
    add_brand(slide, page, dark=dark)
    return slide


def build(dest: Path | None = None) -> Path:
    dest = dest or OUT
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    page = 0

    def P():
        nonlocal page
        page += 1
        return page

    s = new_slide(prs, P(), dark=True)
    add_textbox(s, MX, 2.15, 12, 0.4, "GO PROGRAMMING MASTER COURSE", size=14, bold=True, color=CREAM_TEXT)
    add_textbox(s, MX, 2.55, 12, 1.0, "Week 8, Day 1", size=48, bold=True, color=WHITE)
    add_textbox(s, MX, 3.7, 11, 1.1, "Pointers, memory pictures, and a StudentRepository — HTTP stops owning the map.", size=22, color=RGBColor(0xE7, 0xDD, 0xD0))
    add_textbox(s, MX, 5.3, 11, 0.9, "Umesh Indrajith · BSc Eng (Hons), University of Moratuwa\nSenior Software Engineer / Lecturer", size=16, color=RGBColor(0xD7, 0xCB, 0xB8))
    add_notes(s, "Week 7 create/list must work. Today: pointers + move map into repository. No Postgres. No JWT.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Since last week", "Students live in the handler. That was temporary.", 26)
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "POST / GET / GET:id still work from an in-memory map.",
        "Cold-call: what is a pointer receiver? Why *StudentHandler?",
        "Today we name that idea properly — then extract storage.",
        "Still no database. Repository shape matches this repo; map underneath.",
    ], 18)
    add_notes(s, "Spot-check three create/list demos.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Today", "If Day 1 of week 8 works, you can")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Read and write through & and *; explain a nil pointer panic.",
        "Say why maps/slices already share backing data.",
        "Sketch stack vs heap / GC at picture level.",
        "Move the student map into internal/repository.",
        "Keep handlers talking HTTP only — they call the repo.",
    ], 18)
    add_notes(s, "Lab: crash demo then fix; extract repo.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Software engineering", "HTTP must not own storage forever", 26)
    add_textbox(s, MX, 2.1, 12, 1.0, "A repository hides where data lives. Handlers stay thin: bind JSON, call repo, map results to status codes.", size=18, color=MUTED)
    add_bullets(s, MX, 3.3, 12, 3.2, [
        "Same shape as this repo’s student_repository.go — without pgx yet.",
        "Week 14 swaps the map for SQL; handlers should barely change.",
        "If map writes sit inside a Gin handler, the layering is lying.",
    ], 18)
    add_notes(s, "Point at real repo file.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Pointers", "& address · * value at address")
    add_code(s, MX, 2.05, 12, 2.6, """x := 10
p := &x      // p is *int — points at x
fmt.Println(*p) // 10
*p = 20
fmt.Println(x)  // 20""", 16)
    add_bullets(s, MX, 4.9, 12, 1.8, [
        "A pointer holds a memory address, not the value itself.",
        "Methods that must share or mutate one instance use pointer receivers.",
    ], 18)
    add_notes(s, "Draw box for x and arrow for p.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Nil pointers", "Zero value of a pointer is nil")
    add_code(s, MX, 2.05, 12, 1.8, """var h *handlers.StudentHandler
h.List(c) // panic: nil pointer dereference""", 16)
    add_bullets(s, MX, 4.1, 12, 2.6, [
        "Nil means “no address yet” — not an empty struct.",
        "Always construct with NewStudentHandler(repo) before use.",
        "Lab starts with a deliberate crash, then a one-line fix.",
    ], 18)
    add_notes(s, "Do the crash live before theory overloads.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Pointer receivers", "Why handlers use *StudentHandler", 26)
    add_code(s, MX, 2.05, 12, 2.4, """type StudentHandler struct {
    repo *repository.StudentRepository
}
func (h *StudentHandler) Create(c *gin.Context) { /* uses h.repo */ }""", 15)
    add_bullets(s, MX, 4.7, 12, 2.0, [
        "One handler instance is shared across requests — it holds the repo.",
        "Course default for HTTP handlers: pointer receivers.",
    ], 18)
    add_notes(s, "Connect to week 6 preview.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Value vs reference", "Structs copy · maps and slices share", 26)
    add_card(s, MX, 2.1, 5.9, 1.8, "Struct / int", "Passed by value — callee gets a copy unless you pass a pointer.")
    add_card(s, 7.0, 2.1, 5.7, 1.8, "Map / slice header", "Header copied; backing data shared. Mutations show up.")
    add_card(s, MX, 4.1, 5.9, 1.8, "String", "Immutable bytes; header copied.")
    add_card(s, 7.0, 4.1, 5.7, 1.8, "Pointer", "Explicit shared memory. Check nil before dereference.")
    add_notes(s, "One board sketch.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Stack vs heap · picture", "Where values live (enough for today)", 26)
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Stack — short-lived locals; cheap; gone when the function returns.",
        "Heap — longer-lived / shared; GC cleans unreachable objects.",
        "Returning &local or storing a pointer in a long-lived struct → often heap.",
        "You do not manage free() — the runtime does.",
    ], 18)
    add_notes(s, "Picture level only.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Escape analysis · one slide", "Does this value outlive the function?", 26)
    add_code(s, MX, 2.05, 12, 2.2, """func NewStudentRepository() *StudentRepository {
    return &StudentRepository{students: map[int]models.Student{}, nextID: 1}
}""", 15)
    add_bullets(s, MX, 4.5, 12, 2.2, [
        "Compiler decides stack vs heap. go build -gcflags=\"-m\" is optional.",
        "Rule of thumb: if something else keeps a pointer to it, it escapes.",
    ], 18)
    add_notes(s, "Optional -m flag once if room.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Garbage collection", "Unreachable = reclaimable")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "When nothing points at an object, GC may free it later.",
        "Leaks in Go are usually still-reachable junk (later weeks).",
        "Today’s map store lives until process exit — fine for week 8.",
    ], 20)
    add_notes(s, "One slide.")

    s = new_slide(prs, P())
    add_kicker_title(s, "make vs new · light", "Only if it comes up")
    add_code(s, MX, 2.05, 12, 2.0, """m := make(map[int]models.Student) // ready to use
p := new(StudentRepository)       // *T zero value; maps inside still nil""", 15)
    add_bullets(s, MX, 4.3, 12, 2.4, [
        "Prefer make for maps/slices/channels; constructors for structs.",
        "Course style: NewStudentRepository() that makes the map.",
    ], 18)
    add_notes(s, "Skip live if time is tight.")

    s = new_slide(prs, P())
    add_kicker_title(s, "SMS milestone", "StudentRepository owns the map")
    add_code(s, MX, 2.05, 12, 2.8, """backend/internal/repository/student_repository.go
  Create / GetAll / GetByID (/ Update / Delete)

handlers.StudentHandler { repo *repository.StudentRepository }
// handlers call repo — no map fields on the handler""", 14)
    add_bullets(s, MX, 5.1, 12, 1.6, [
        "Routes and status codes stay the same as week 7.",
        "No JWT. No Postgres. Restart still clears memory.",
    ], 18)
    add_notes(s, "Board the package split.")

    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "Nil crash", "Break it on purpose", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 2.6, """var h *handlers.StudentHandler
r.GET("/api/students", h.List) // registers fine…
// first request → panic: invalid memory address""", 15)
    add_bullets(s, MX, 5.1, 12, 1.6, [
        "Show the stack trace. Name the bug: nil receiver.",
        "Fix: h := handlers.NewStudentHandler(repo) before wiring.",
    ], 18)
    add_notes(s, "Students must see the panic once.")

    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "Repository", "Extract the map", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 4.0, """type StudentRepository struct {
    mu       sync.Mutex
    students map[int]models.Student
    nextID   int
}
func NewStudentRepository() *StudentRepository {
    return &StudentRepository{students: map[int]models.Student{}, nextID: 1}
}""", 14)
    add_notes(s, "Move Create/List/GetByID bodies into repo methods.")

    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "Handler", "Thin HTTP over the repo", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 4.0, """func NewStudentHandler(repo *repository.StudentRepository) *StudentHandler {
    return &StudentHandler{repo: repo}
}
func (h *StudentHandler) GetByID(c *gin.Context) {
    s, ok := h.repo.GetByID(id)
    if !ok { c.JSON(404, gin.H{"error": "student not found"}); return }
    c.JSON(200, s)
}""", 14)
    add_notes(s, "curl create + list after extract.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Wiring", "What main looks like after")
    add_code(s, MX, 2.05, 12, 2.6, """repo := repository.NewStudentRepository()
students := handlers.NewStudentHandler(repo)
r.POST("/api/students", students.Create)
r.GET("/api/students", students.List)
r.GET("/api/students/:id", students.GetByID)""", 15)
    add_bullets(s, MX, 4.9, 12, 1.8, [
        "Construct repo first, inject into handler — DI without a framework.",
        "Later weeks inject DB pool the same way.",
    ], 18)
    add_notes(s, "Contrast with week 7 NewStudentHandler() that owned the map.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Not found", "Repo returns data · handler picks status", 26)
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Repo: (Student, bool) or (Student, error) — pick one style and stick to it.",
        "Handler maps missing → 404, bad id → 400, success → 200.",
        "Do not import gin inside repository.",
    ], 20)
    add_notes(s, "Hard rule: repository package has no HTTP.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Lab · rest of class", "Done when all of this is true")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Nil-pointer crash demonstrated and fixed",
        "internal/repository/student_repository.go owns the map",
        "Handler holds *StudentRepository; no map fields on handler",
        "Postman create + list + get still pass; health green",
        "PR started describing the move",
    ], 18)
    add_notes(s, "Gate on board.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Homework", "PR the extract + draw memory")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Handler → repository for create/list/get (and update/delete if present).",
        "PR description: notebook photo or ASCII of stack/heap for &x / nil crash.",
        "Still no JWT / Postgres.",
        "Stretch: repo returns error; ErrNotFound → 404 (week 9 preview).",
    ], 18)
    add_notes(s, "Accept phone photos of whiteboard.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Next session", "Week 9 — errors, validation, slog, JWT preview", 24)
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "fmt.Errorf with %w, errors.Is / As.",
        'Consistent {"error":"..."}; structured logs.',
        "Bring a clean handler → repository split.",
    ], 20)
    add_notes(s, "Do not teach wrapping today beyond a teaser.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Words to keep", "Say these correctly next week")
    add_card(s, MX, 2.1, 5.9, 1.8, "Pointer", "Address of a value; nil means none.")
    add_card(s, 7.0, 2.1, 5.7, 1.8, "Receiver", "Method’s target; pointer shares one instance.")
    add_card(s, MX, 4.1, 5.9, 1.8, "Repository", "Storage API; hides map or SQL.")
    add_card(s, 7.0, 4.1, 5.7, 1.8, "Escape", "Value outlives the function → often heap.")
    add_notes(s, "Oral quiz.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Layout reminder", "Where the new files live")
    add_code(s, MX, 2.05, 12, 4.0, """backend/
  cmd/server/main.go
  internal/models/student.go
  internal/repository/
    student_repository.go
  internal/handlers/
    student.go                    # HTTP only; holds *StudentRepository""", 15)
    add_notes(s, "Ignore database package until week 14.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Honest limits", "What this week is not")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Not Postgres — repository still wraps a map.",
        "Not a deep OS memory course — pictures beat dumps.",
        "Not interface-heavy DI — concrete *StudentRepository is fine.",
        "Races under load still week 12.",
    ], 18)
    add_notes(s, "Keep spiral honest.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Safety", "Still true while refactoring")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Commit working week-7 routes before the extract.",
        "Do not “fix” by deleting go.mod.",
        "gofmt; no secrets; port :8080.",
    ], 20)
    add_notes(s, "Small steps: crash fix commit, then repo extract commit.")

    s = new_slide(prs, P(), dark=True)
    add_textbox(s, MX, 3.0, 12, 1.2, "Questions", size=64, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_notes(s, "Then nil demo → extract repo → lab.")

    if page != TOTAL:
        raise SystemExit(f"expected {TOTAL} slides, built {page}")
    prs.save(str(dest))
    print("Wrote", dest)
    return dest


if __name__ == "__main__":
    build()
