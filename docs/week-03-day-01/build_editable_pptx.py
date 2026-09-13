#!/usr/bin/env python3
"""Build an editable PowerPoint of the Week 3 Day 1 deck (native text, not screenshots)."""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    req = Path(__file__).resolve().parent / "requirements.txt"
    sys.stderr.write(
        "Missing python-pptx. From the repo root run:\n"
        f"  python3 -m pip install --user -r {req}\n"
    )
    raise

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[1]
LOGO = REPO / "docs" / "branding" / "epic-learn-logo.png"
OUT = ROOT / "pdf" / "week-03-day-01-slides.pptx"
FOOTER = "Epic Learn Institute of Higher Education - Go Programming Master Course"

INK = RGBColor(0x14, 0x20, 0x2B)
MUTED = RGBColor(0x4D, 0x5D, 0x6B)
TEAL = RGBColor(0x0B, 0x6E, 0x68)
TEAL_DARK = RGBColor(0x08, 0x4F, 0x4B)
NAVY = RGBColor(0x0A, 0x2A, 0x5C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF4, 0xEF, 0xE6)
CARD = RGBColor(0xFF, 0xFD, 0xF9)
LINE = RGBColor(0xD8, 0xCF, 0xC2)
CODE_BG = RGBColor(0x10, 0x21, 0x2C)
CODE_FG = RGBColor(0xE8, 0xF2, 0xEF)
GOLD = RGBColor(0xD9, 0xB8, 0x26)
TITLE_BG = RGBColor(0x14, 0x2A, 0x32)
CREAM_TEXT = RGBColor(0xEA, 0xD9, 0xA3)
LIVE = RGBColor(0xC2, 0x4E, 0x1D)

W, H = 13.333, 7.5
MX = 0.65
HEADER_H = 0.78
FOOTER_H = 0.42
CONTENT_TOP = 0.95
TOTAL = 29


def _font(run, name="Calibri", size=18, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_textbox(slide, l, t, w, h, text, size=18, bold=False, color=INK, font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _font(run, font, size, bold, color)
    return box


def add_bullets(slide, l, t, w, h, items, size=20, color=INK):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        p.level = 0
        run = p.add_run()
        run.text = "•  " + item
        _font(run, "Calibri", size, False, color)
    return box


def add_code(slide, l, t, w, h, text, size=15):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = CODE_BG
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.1)
    lines = text.strip("\n").split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line if line else " "
        _font(run, "Consolas", size, False, CODE_FG)
    return shape


def add_card(slide, l, t, w, h, title, body, title_size=16, body_size=13):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = LINE
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    _font(run, "Calibri", title_size, True, TEAL_DARK)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    run2 = p2.add_run()
    run2.text = body
    _font(run2, "Calibri", body_size, False, INK)
    return shape


def add_brand(slide, page, dark=False):
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(HEADER_H))
    _fill(header, WHITE)
    slide.shapes.add_picture(str(LOGO), Inches(0.35), Inches(0.12), height=Inches(0.55))
    add_textbox(
        slide, 10.2, 0.22, 2.8, 0.4,
        "WEEK 3  ·  DAY 1",
        size=12, bold=True, color=CREAM_TEXT if dark else TEAL_DARK, align=PP_ALIGN.RIGHT,
    )
    gold = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(HEADER_H), Inches(W - 0.7), Inches(0.02)
    )
    _fill(gold, GOLD)

    footer_bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(H - FOOTER_H), Inches(W), Inches(FOOTER_H)
    )
    _fill(footer_bg, WHITE)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(H - FOOTER_H), Inches(W - 0.7), Inches(0.015)
    )
    _fill(line, NAVY)
    add_textbox(
        slide, 0.5, H - FOOTER_H + 0.06, 10.5, 0.3,
        FOOTER, size=11, color=NAVY, align=PP_ALIGN.LEFT,
    )
    add_textbox(
        slide, 11.3, H - FOOTER_H + 0.06, 1.6, 0.3,
        f"{page} / {TOTAL}", size=11, color=MUTED, align=PP_ALIGN.RIGHT,
    )


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_kicker_title(slide, kicker, title, title_size=32, top=None):
    top = CONTENT_TOP if top is None else top
    add_textbox(slide, MX, top, 12, 0.32, kicker.upper(), size=13, bold=True, color=TEAL)
    add_textbox(slide, MX, top + 0.28, 12, 0.7, title, size=title_size, bold=True, color=INK)


def new_slide(prs, page, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if dark:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
        _fill(bg, TITLE_BG)
    else:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
        _fill(bg, CREAM)
    add_brand(slide, page, dark=dark)
    return slide




def build(dest: Path | None = None) -> Path:
    dest = dest or OUT
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    page = 0

    def P():
        nonlocal page
        page += 1
        return page

    # 1 title
    s = new_slide(prs, P(), dark=True)
    add_textbox(s, MX, 2.15, 12, 0.4, "GO PROGRAMMING MASTER COURSE", size=14, bold=True, color=CREAM_TEXT)
    add_textbox(s, MX, 2.55, 12, 1.0, "Week 3, Day 1", size=48, bold=True, color=WHITE)
    add_textbox(
        s, MX, 3.7, 11, 1.1,
        "Control flow, slices, maps, and just enough functions — the tools you will use to filter students later.",
        size=22, color=RGBColor(0xE7, 0xDD, 0xD0),
    )
    add_textbox(
        s, MX, 5.3, 11, 0.9,
        "Umesh Indrajith · BSc Eng (Hons), University of Moratuwa\nSenior Software Engineer / Lecturer",
        size=16, color=RGBColor(0xD7, 0xCB, 0xB8),
    )
    add_notes(s, "Welcome. First 10–15 min: green FullName tests. Dense week — two labs. Still no SMS. SMS starts week 4.")

    # 2 recap
    s = new_slide(prs, P())
    add_kicker_title(s, "Since last week", "Types were the foundation. Prove it.")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "go test on FullName is green — including trim cases. If not, fix that first.",
        "Cold-call: zero value of string? Why is len(\"café\") 5?",
        "gofmt still required. Unformatted PRs are returned.",
        "Today is dense. We will do two labs: control+slices, then maps.",
    ], 20)
    add_notes(s, "Spot-check three repos. Do not start slices on broken tests.")

    # 3 outcomes
    s = new_slide(prs, P())
    add_kicker_title(s, "Today", "If Day 1 of week 3 works, you can")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Write if, switch, and every shape of for / range.",
        "Tell an array from a slice; use make, append, len, cap.",
        "Explain the backing-array gotcha after append.",
        "Read and write maps with the comma-ok check.",
        "Write a function that returns (value, error) as a pattern you copy.",
        "Ship Contains, Unique, and a frequency counter with green go test.",
    ], 18)
    add_notes(s, "Prefer two labs over rushing maps. Functions today are survival kit for week 5.")

    # 4 SE
    s = new_slide(prs, P())
    add_kicker_title(s, "Software engineering", "Control flow is how requirements become behaviour", 26)
    add_textbox(
        s, MX, 2.1, 12, 1.2,
        "A requirement says “reject empty email” or “filter by course”. That is an if, a loop, and a collection — not a framework.",
        size=18, color=MUTED,
    )
    add_bullets(s, MX, 3.4, 12, 3.0, [
        "SMS later: list students, keep those whose course matches, count by course.",
        "Today’s katas become those utilities. Write tests on paper first.",
        "Still no HTTP. Still no database. Language that ships product.",
    ], 18)
    add_notes(s, "Draw Contains on the board with 3 rows before they open the editor.")

    # 5 if
    s = new_slide(prs, P())
    add_kicker_title(s, "if", "Conditions are boolean — nothing else")
    add_code(s, MX, 2.05, 12, 2.6, """score := 72
if score >= 50 {
    fmt.Println("pass")
} else if score >= 40 {
    fmt.Println("resit")
} else {
    fmt.Println("fail")
}""", 16)
    add_bullets(s, MX, 4.8, 12, 1.8, [
        "No parentheses around the condition. Braces are required — even for one line.",
        "There is no if score. An int is not a bool.",
    ], 18)
    add_notes(s, "Students from C/JS will write if (x) or omit braces.")

    # 6 short if
    s = new_slide(prs, P())
    add_kicker_title(s, "if", "Short statement — declare, then decide")
    add_code(s, MX, 2.05, 12, 2.4, """if n := len(name); n == 0 {
    fmt.Println("empty")
} else {
    fmt.Println("chars:", n)
}
// n is not visible here""", 16)
    add_bullets(s, MX, 4.6, 12, 2.0, [
        "The short statement runs first. The name lives only inside the if/else block.",
        "You will see if err := …; err != nil from week 5 on. Learn the shape now.",
    ], 18)
    add_notes(s, "Plant err != nil without teaching error theory.")

    # 7 switch
    s = new_slide(prs, P())
    add_kicker_title(s, "switch", "Many equal cases, one place")
    add_code(s, MX, 2.05, 12, 3.0, """day := "Sat"
switch day {
case "Sat", "Sun":
    fmt.Println("weekend")
case "Mon", "Tue", "Wed", "Thu", "Fri":
    fmt.Println("weekday")
default:
    fmt.Println("unknown")
}""", 15)
    add_bullets(s, MX, 5.2, 12, 1.5, [
        "No automatic fallthrough. Cases can list several values with commas.",
    ], 18)
    add_notes(s, "Contrast with C/Java fallthrough.")

    # 8 tagless
    s = new_slide(prs, P())
    add_kicker_title(s, "switch", "Tagless switch — a clean if/else if chain")
    add_code(s, MX, 2.05, 12, 3.2, """score := 72
switch {
case score >= 75:
    fmt.Println("A")
case score >= 65:
    fmt.Println("B")
case score >= 50:
    fmt.Println("C")
default:
    fmt.Println("F")
}""", 15)
    add_bullets(s, MX, 5.4, 12, 1.3, [
        "No expression after switch. Cases are boolean conditions.",
    ], 18)
    add_notes(s, "Useful for validation rules later.")

    # 9 for
    s = new_slide(prs, P())
    add_kicker_title(s, "for", "Go has one loop keyword")
    add_card(s, MX, 2.1, 5.9, 1.8, "C-style", "for i := 0; i < 5; i++ { … }")
    add_card(s, 7.0, 2.1, 5.7, 1.8, "While-style", "for n > 0 { n--; … }")
    add_card(s, MX, 4.1, 5.9, 1.8, "Forever", "for { break }")
    add_card(s, 7.0, 4.1, 5.7, 1.8, "No while / do", "Those words do not exist. for covers every shape.")
    add_notes(s, "Write all three on the board.")

    # 10 range
    s = new_slide(prs, P())
    add_kicker_title(s, "range", "Walk a collection without inventing an index")
    add_code(s, MX, 2.05, 12, 2.8, """courses := []string{"SE", "CS", "IT"}
for i, c := range courses {
    fmt.Println(i, c)
}
for _, c := range courses {
    fmt.Println(c)
}""", 15)
    add_bullets(s, MX, 5.0, 12, 1.6, [
        "_ means “I know there is a value; I am not using it.”",
        "On a string, range yields runes (you saw this in week 2).",
    ], 18)
    add_notes(s, "Blank identifier is required — unused i is a compile error.")

    # 11 break/continue
    s = new_slide(prs, P())
    add_kicker_title(s, "break / continue", "Leave early, or skip this turn")
    add_code(s, MX, 2.05, 12, 2.8, """for i := 0; i < 10; i++ {
    if i%2 == 0 {
        continue
    }
    if i > 7 {
        break
    }
    fmt.Println(i)
}""", 15)
    add_bullets(s, MX, 5.0, 12, 1.6, [
        "continue → next iteration. break → exit the innermost loop.",
        "Labels exist for nested loops. You do not need them today.",
    ], 18)
    add_notes(s, "Quick demo. Do not teach labeled break.")

    # 12 live control
    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "Control flow", "Type this with me", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 3.8, """package main
import "fmt"
func main() {
    grade := "B"
    switch grade {
    case "A", "B":
        fmt.Println("good")
    default:
        fmt.Println("work")
    }
    for i, r := range "Go" {
        fmt.Printf("%d %c\\n", i, r)
    }
}""", 15)
    add_notes(s, "File: ~/epic-go/week03/control.go")

    # 13 arrays vs slices
    s = new_slide(prs, P())
    add_kicker_title(s, "Arrays vs slices", "Fixed size vs the list you will actually use", 28)
    add_bullets(s, MX, 2.1, 6.5, 3.5, [
        "An array has a fixed length in its type: [3]int.",
        "A slice is a view: pointer + length + capacity. Type is []int.",
        "Almost all Go code uses slices.",
    ], 18)
    add_code(s, 7.3, 2.2, 5.3, 2.6, """var seats [3]int
courses := []string{"SE"}
more := make([]string, 0, 8)""", 14)
    add_notes(s, "Draw ptr, len, cap. [3]int and [4]int are different types.")

    # 14 len/cap/make
    s = new_slide(prs, P())
    add_kicker_title(s, "len / cap / make", "How big is it, and how big can it grow?", 28)
    add_code(s, MX, 2.05, 12, 2.2, """s := make([]int, 2, 5) // len 2, cap 5
fmt.Println(len(s), cap(s), s) // 2 5 [0 0]
s = append(s, 9)
fmt.Println(len(s), cap(s), s) // 3 5 [0 0 9]""", 15)
    add_bullets(s, MX, 4.5, 12, 2.2, [
        "len — how many elements you can index.",
        "cap — size of the backing array from the start of this slice.",
        "make([]T, length, capacity) — prefer when you will append a lot.",
    ], 18)
    add_notes(s, "Cap sets up the gotcha.")

    # 15 append
    s = new_slide(prs, P())
    add_kicker_title(s, "append", "Always take the result")
    add_code(s, MX, 2.05, 12, 2.2, """names := []string{"Ada"}
names = append(names, "Grace")
names = append(names, "Edsger", "Ken")
// wrong: append(names, "x")  — threw away the new header""", 15)
    add_bullets(s, MX, 4.5, 12, 2.2, [
        "append may reuse the backing array or allocate a bigger one.",
        "Either way, keep the returned slice header.",
        "Appending to a nil slice is fine.",
    ], 18)
    add_notes(s, "Show nil append: var xs []int; xs = append(xs, 1).")

    # 16 gotcha
    s = new_slide(prs, P())
    add_kicker_title(s, "The gotcha", "Two slices can share a backing array")
    add_code(s, MX, 2.05, 12, 2.2, """a := []int{10, 20, 30, 40}
b := a[:2]
b = append(b, 99)
fmt.Println(a) // often [10 20 99 40] — surprise""", 15)
    add_bullets(s, MX, 4.5, 12, 2.2, [
        "If cap(b) still has room, append writes into the shared array.",
        "Need independence: append([]int(nil), a...) or copy.",
        "This is the week-3 idea they must take home — same weight as zero values.",
    ], 18)
    add_notes(s, "LIVE this. Draw the array boxes. Do not skip.")

    # 17 live slices
    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "Slices", "Watch the shared array", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 3.6, """package main
import "fmt"
func main() {
    a := []int{10, 20, 30, 40}
    b := a[:2]
    fmt.Println("before", a, b, "cap(b)=", cap(b))
    b = append(b, 99)
    fmt.Println("after ", a, b)
}""", 15)
    add_notes(s, "File: ~/epic-go/week03/slices.go. Then show full-cap allocate case.")

    # 18 maps
    s = new_slide(prs, P())
    add_kicker_title(s, "maps", "Lookup by key — the dictionary type")
    add_code(s, MX, 2.05, 12, 2.4, """ages := map[string]int{
    "Ada": 36,
    "Grace": 85,
}
ages["Ken"] = 70
fmt.Println(ages["Ada"])
delete(ages, "Grace")""", 15)
    add_bullets(s, MX, 4.7, 12, 2.0, [
        "Type is map[Key]Value. Keys must be comparable.",
        "Zero value is nil. Reading returns zero; writing panics — make first.",
    ], 18)
    add_notes(s, "Demo nil map write panic once.")

    # 19 comma-ok
    s = new_slide(prs, P())
    add_kicker_title(s, "comma-ok", "Missing key is not an error — check yourself")
    add_code(s, MX, 2.05, 12, 2.4, """v, ok := ages["Ada"]
if !ok {
    fmt.Println("not found")
} else {
    fmt.Println(v)
}
fmt.Println(ages["Missing"]) // 0 — zero value, not a crash""", 15)
    add_bullets(s, MX, 4.7, 12, 2.0, [
        "v := ages[\"x\"] cannot tell missing from present-with-zero.",
        "Comma-ok is the idiom. Memorise it.",
    ], 18)
    add_notes(s, "Drill it. Map equivalent of zero-value literacy.")

    # 20 range map
    s = new_slide(prs, P())
    add_kicker_title(s, "range over map", "Order is not promised")
    add_code(s, MX, 2.05, 12, 2.0, """for name, age := range ages {
    fmt.Println(name, age)
}""", 16)
    add_bullets(s, MX, 4.3, 12, 2.4, [
        "Iteration order is randomised. Never require a fixed map print order in tests.",
        "Need sorted keys? Collect keys into a slice, sort.Strings, then look up.",
    ], 18)
    add_notes(s, "Run range twice — order may differ.")

    # 21 live freq
    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "Maps", "Frequency counter — lab 2 preview", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 3.8, """words := []string{"SE", "CS", "SE", "IT", "SE"}
count := make(map[string]int)
for _, w := range words {
    count[w]++
}
fmt.Println(count)
n, ok := count["SE"]
fmt.Println(n, ok)""", 15)
    add_notes(s, "count[w]++ works because missing key reads as 0.")

    # 22 functions
    s = new_slide(prs, P())
    add_kicker_title(s, "Functions · survival kit", "Enough to survive week 5 — not full theory", 26)
    add_code(s, MX, 2.0, 12, 3.4, """func Contains(list []string, want string) bool {
    for _, v := range list {
        if v == want { return true }
    }
    return false
}
func FullName(first, last string) (string, error) {
    if first == "" { return "", fmt.Errorf("first required") }
    return first + " " + last, nil
}""", 14)
    add_bullets(s, MX, 5.5, 12, 1.2, [
        "return x, err is a pattern you copy. Deep error theory is week 9.",
    ], 18)
    add_notes(s, "Week 6 deepens functions. No defer/methods today.")

    # 23 variadic light
    s = new_slide(prs, P())
    add_kicker_title(s, "Functions · light", "Variadic and closures — know they exist", 28)
    add_code(s, MX, 2.05, 12, 2.8, """func Sum(nums ...int) int {
    total := 0
    for _, n := range nums { total += n }
    return total
}
double := func(n int) int { return n * 2 }""", 15)
    add_bullets(s, MX, 5.1, 12, 1.6, [
        "... means any number of args. Inside, it is a slice.",
        "Pace valve: skip live typing if slices/maps ran long.",
    ], 18)
    add_notes(s, "Show the slide; skip typing if late.")

    # 24 test table
    s = new_slide(prs, P())
    add_kicker_title(s, "Software engineering", "Write the test table before the code", 28)
    add_code(s, MX, 2.05, 12, 2.2, """Input list          want     result
["SE","CS"]         "SE"     true
["SE","CS"]         "IT"     false
[]                  "SE"     false""", 15)
    add_bullets(s, MX, 4.5, 12, 2.2, [
        "Then implement. Then go test. Red → green is still the ritual.",
        "Same habit later for every SMS handler edge case.",
    ], 18)
    add_notes(s, "Force paper tables in lab 1.")

    # 25 lab 1
    s = new_slide(prs, P())
    add_kicker_title(s, "Lab 1 · control + slices", "Contains and Unique")
    add_bullets(s, MX, 2.15, 12, 2.5, [
        "Package slicekit with green tests.",
        "Contains(list []string, want string) bool",
        "Unique(list []string) []string — first occurrence wins, order preserved.",
        "Fill the paper test table first. Then code.",
    ], 18)
    add_code(s, MX, 4.7, 12, 1.8, """mkdir -p ~/epic-go/week03/slicekit
cd ~/epic-go/week03/slicekit && go mod init epiclearn/week03
go test""", 14)
    add_notes(s, "Unique needs a seen map + append to result.")

    # 26 lab 2
    s = new_slide(prs, P())
    add_kicker_title(s, "Lab 2 · maps", "Frequency counter with tests")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Freq(words []string) map[string]int",
        "Empty input → empty map (not nil panic).",
        "Assert counts with comma-ok — do not compare map print order.",
        "Both labs must be green before you leave.",
    ], 20)
    add_notes(s, "If time is short, Freq stays in the same package.")

    # 27 gate
    s = new_slide(prs, P())
    add_kicker_title(s, "Lab gate", "Done when all of this is true")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Worksheet: control-flow + slice gotcha + comma-ok cells filled",
        "Paper test table for Contains completed before coding",
        "go test green for Contains, Unique, and Freq",
        "You can explain why append on a subslice changed the original",
        "Commit and push. Show me the GitHub URL",
    ], 20)
    add_notes(s, "Checklist on the board. Pair done with stuck.")

    # 28 next
    s = new_slide(prs, P())
    add_kicker_title(s, "Next session", "Week 4 — modules, packages, SMS kickoff")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "go mod init with meaning. internal/ as a boundary.",
        "Project layout matching this repo’s backend.",
        "First SMS milestone: empty module + folders + README.",
        "Bring green week-3 tests. Product work starts.",
    ], 20)
    add_notes(s, "Do not start modules theory today.")

    # 29 questions
    s = new_slide(prs, P(), dark=True)
    add_textbox(s, MX, 3.0, 12, 1.2, "Questions", size=64, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_notes(s, "Take questions. Start lab 1. Collect URLs.")

    if page != TOTAL:
        raise SystemExit(f"expected {TOTAL} slides, built {page}")

    prs.save(str(dest))
    print("Wrote", dest)
    return dest


if __name__ == "__main__":
    build()
