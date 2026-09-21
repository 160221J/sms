#!/usr/bin/env python3
"""Build an editable PowerPoint of the Week 5 Day 1 deck (native text, not screenshots)."""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    req = Path(__file__).resolve().parent / "requirements.txt"
    sys.stderr.write(
        "Missing python-pptx. From the repo root run:\n"
        f"  python3 -m pip install --user -r {req}\n"
    )
    raise

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[1]
LOGO = REPO / "docs" / "branding" / "epic-learn-logo.png"
OUT = ROOT / "pdf" / "week-05-day-01-slides.pptx"
FOOTER = "Epic Learn Institute of Higher Education - Go Programming Master Course"

INK = RGBColor(0x14, 0x20, 0x2B)
MUTED = RGBColor(0x4D, 0x5D, 0x6B)
TEAL = RGBColor(0x0B, 0x6E, 0x68)
TEAL_DARK = RGBColor(0x08, 0x4F, 0x4B)
NAVY = RGBColor(0x0A, 0x2A, 0x5C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF4, 0xEF, 0xE6)
CARD = RGBColor(0xFF, 0xFD, 0xF9)
LINE = RGBColor(0xD8, 0xCF, 0xC2)
CODE_BG = RGBColor(0x10, 0x21, 0x2C)
CODE_FG = RGBColor(0xE8, 0xF2, 0xEF)
GOLD = RGBColor(0xD9, 0xB8, 0x26)
TITLE_BG = RGBColor(0x14, 0x2A, 0x32)
CREAM_TEXT = RGBColor(0xEA, 0xD9, 0xA3)
LIVE = RGBColor(0xC2, 0x4E, 0x1D)

W, H = 13.333, 7.5
MX = 0.65
HEADER_H = 0.78
FOOTER_H = 0.42
CONTENT_TOP = 0.95
TOTAL = 25


def _font(run, name="Calibri", size=18, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_textbox(slide, l, t, w, h, text, size=18, bold=False, color=INK, font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _font(run, font, size, bold, color)
    return box


def add_bullets(slide, l, t, w, h, items, size=20, color=INK):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        p.level = 0
        run = p.add_run()
        run.text = "•  " + item
        _font(run, "Calibri", size, False, color)
    return box


def add_code(slide, l, t, w, h, text, size=15):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = CODE_BG
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.1)
    lines = text.strip("\n").split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line if line else " "
        _font(run, "Consolas", size, False, CODE_FG)
    return shape


def add_card(slide, l, t, w, h, title, body, title_size=16, body_size=13):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = LINE
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    _font(run, "Calibri", title_size, True, TEAL_DARK)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    run2 = p2.add_run()
    run2.text = body
    _font(run2, "Calibri", body_size, False, INK)
    return shape


def add_brand(slide, page, dark=False):
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(HEADER_H))
    _fill(header, WHITE)
    slide.shapes.add_picture(str(LOGO), Inches(0.35), Inches(0.12), height=Inches(0.55))
    add_textbox(
        slide, 10.2, 0.22, 2.8, 0.4,
        "WEEK 5  ·  DAY 1",
        size=12, bold=True, color=CREAM_TEXT if dark else TEAL_DARK, align=PP_ALIGN.RIGHT,
    )
    gold = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(HEADER_H), Inches(W - 0.7), Inches(0.02)
    )
    _fill(gold, GOLD)

    footer_bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(H - FOOTER_H), Inches(W), Inches(FOOTER_H)
    )
    _fill(footer_bg, WHITE)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(H - FOOTER_H), Inches(W - 0.7), Inches(0.015)
    )
    _fill(line, NAVY)
    add_textbox(
        slide, 0.5, H - FOOTER_H + 0.06, 10.5, 0.3,
        FOOTER, size=11, color=NAVY, align=PP_ALIGN.LEFT,
    )
    add_textbox(
        slide, 11.3, H - FOOTER_H + 0.06, 1.6, 0.3,
        f"{page} / {TOTAL}", size=11, color=MUTED, align=PP_ALIGN.RIGHT,
    )


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_kicker_title(slide, kicker, title, title_size=32, top=None):
    top = CONTENT_TOP if top is None else top
    add_textbox(slide, MX, top, 12, 0.32, kicker.upper(), size=13, bold=True, color=TEAL)
    add_textbox(slide, MX, top + 0.28, 12, 0.7, title, size=title_size, bold=True, color=INK)


def new_slide(prs, page, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if dark:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
        _fill(bg, TITLE_BG)
    else:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
        _fill(bg, CREAM)
    add_brand(slide, page, dark=dark)
    return slide


def build(dest: Path | None = None) -> Path:
    dest = dest or OUT
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    page = 0

    def P():
        nonlocal page
        page += 1
        return page

    # 1 — Title
    s = new_slide(prs, P(), dark=True)
    add_textbox(s, MX, 2.15, 12, 0.4, "GO PROGRAMMING MASTER COURSE", size=14, bold=True, color=CREAM_TEXT)
    add_textbox(s, MX, 2.55, 12, 1.0, "Week 5, Day 1", size=48, bold=True, color=WHITE)
    add_textbox(s, MX, 3.7, 11, 1.1, "HTTP, JSON, and Gin — the first SMS API you can hit from Postman.", size=22, color=RGBColor(0xE7, 0xDD, 0xD0))
    add_textbox(s, MX, 5.3, 11, 0.9, "Umesh Indrajith · BSc Eng (Hons), University of Moratuwa\nSenior Software Engineer / Lecturer", size=16, color=RGBColor(0xD7, 0xCB, 0xB8))
    add_notes(s, "Product week 2. Skeleton from week 4 must exist. Today: net/http first, then Gin, GET /health on :8080. No DB. No JWT.")

    # 2 — Since last week
    s = new_slide(prs, P())
    add_kicker_title(s, "Since last week", "Skeleton first. Then the network.")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "go run ./cmd/server still prints sms starting (or we fix that in 10 minutes).",
        "Module path and internal/ tree are in place. PR from week 4 is reviewable.",
        "Today that process learns to speak HTTP.",
        "gofmt and no .env in git — unchanged rules.",
    ], 20)
    add_notes(s, "Spot-check three skeletons. Do not start Gin on a broken module.")

    # 3 — Outcomes
    s = new_slide(prs, P())
    add_kicker_title(s, "Today", "If Day 1 of week 5 works, you can")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Name HTTP methods, paths, and common status codes.",
        "Write a tiny net/http server so Gin is not magic.",
        "Use Gin: router, gin.Context, c.JSON.",
        "Return JSON with a minimal struct + tags.",
        "Ship GET /health on :8080 and prove it in Postman + curl.",
    ], 18)
    add_notes(s, "Deep functions week 6; deep structs week 7. Handlers are functions the router calls.")

    # 4 — Why REST before full functions
    s = new_slide(prs, P())
    add_kicker_title(s, "Software engineering", "Why REST before a full functions course?", 28)
    add_textbox(s, MX, 2.1, 12, 1.2, "So every remaining week has a running product. You copy a handler shape now; you deepen functions and structs inside SMS next.", size=18, color=MUTED)
    add_bullets(s, MX, 3.4, 12, 3.0, [
        "Spiral, not waterfall — just enough to ship a vertical slice.",
        "Today’s milestone: health JSON. Not login. Not Postgres.",
        "A classmate should hit your API from Postman without reading your brain.",
    ], 18)
    add_notes(s, "Reassure: incomplete theory is intentional. Product clock matters.")

    # 5 — HTTP method/path/status/body cards
    s = new_slide(prs, P())
    add_kicker_title(s, "HTTP", "Client asks. Server answers.")
    add_card(s, MX, 2.1, 5.9, 1.8, "Method", "GET read · POST create · PUT replace · DELETE remove")
    add_card(s, 7.0, 2.1, 5.7, 1.8, "Path", "/health · /api/hello · later /api/students")
    add_card(s, MX, 4.1, 5.9, 1.8, "Status", "200 OK · 400 bad input · 404 missing · 500 server fault")
    add_card(s, 7.0, 4.1, 5.7, 1.8, "Body", "Often JSON. Headers carry metadata (Content-Type).")
    add_notes(s, "Board these four words. Today is almost all GET + 200.")

    # 6 — JSON health example
    s = new_slide(prs, P())
    add_kicker_title(s, "JSON", "The body shape APIs speak")
    add_code(s, MX, 2.05, 12, 2.0, """{
  "status": "ok",
  "message": "Backend is running"
}""", 16)
    add_bullets(s, MX, 4.3, 12, 2.4, [
        "Keys are strings. Values: string, number, bool, object, array, null.",
        "Go turns structs ↔ JSON with encoding/json (and Gin helpers).",
        "Wrong JSON from a client → 400 later. Today we mostly send JSON.",
    ], 18)
    add_notes(s, "Show Postman pretty JSON. Match COURSE.md health payload exactly.")

    # 7 — Live net/http hello
    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "stdlib first", "net/http Hello — 30 minutes, no Gin", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 3.2, """package main
import ("fmt"; "net/http")
func main() {
    http.HandleFunc("/hello", func(w http.ResponseWriter, r *http.Request) {
        fmt.Fprintln(w, "hello from stdlib")
    })
    http.ListenAndServe(":8080", nil)
}""", 14)
    add_textbox(s, MX, 5.6, 12, 0.6, "curl http://localhost:8080/hello — Gin is a nicer router on top of this idea.", size=16, color=MUTED)
    add_notes(s, "Separate throwaway file or branch. Students must see ListenAndServe before Gin Default().")

    # 8 — Gin intro + go get
    s = new_slide(prs, P())
    add_kicker_title(s, "Gin", "A router + helpers on top of net/http", 28)
    add_bullets(s, MX, 2.15, 12, 2.4, [
        "Framework rule for this course: stdlib first, Gin for SMS.",
        "You must be able to say what Gin is wrapping — routing, JSON, middleware later.",
        "Add the dependency once you are in backend/:",
    ], 18)
    add_code(s, MX, 4.7, 12, 1.4, """go get github.com/gin-gonic/gin
go mod tidy""", 16)
    add_notes(s, "Live go get. Show go.mod require appear. Windows/WSL: stay in Ubuntu.")

    # 9 — Router context c.JSON
    s = new_slide(prs, P())
    add_kicker_title(s, "Gin", "Router, context, JSON response")
    add_code(s, MX, 2.05, 12, 2.8, """r := gin.Default()
r.GET("/health", func(c *gin.Context) {
    c.JSON(200, gin.H{
        "status":  "ok",
        "message": "Backend is running",
    })
})
r.Run(":8080")""", 15)
    add_bullets(s, MX, 5.1, 12, 1.6, [
        "gin.Context is the per-request bag: params, body, response writers.",
        "gin.H is map[string]any — fine for health. Prefer structs as APIs grow.",
    ], 17)
    add_notes(s, "Prefer http.StatusOK over magic 200 when you show the polished version.")

    # 10 — Minimal struct json tags
    s = new_slide(prs, P())
    add_kicker_title(s, "Structs · minimal", "Just enough for JSON tags", 28)
    add_code(s, MX, 2.05, 12, 3.0, """type HealthResponse struct {
    Status  string `json:"status"`
    Message string `json:"message"`
}

c.JSON(http.StatusOK, HealthResponse{
    Status:  "ok",
    Message: "Backend is running",
})""", 14)
    add_bullets(s, MX, 5.3, 12, 1.5, [
        "Exported fields (capital letters) are what JSON can see.",
        "Tag renames the wire name. Deep structs are week 7 — copy this pattern today.",
    ], 17)
    add_notes(s, "One struct only. Do not teach embedding or methods.")

    # 11 — GET /health contract
    s = new_slide(prs, P())
    add_kicker_title(s, "SMS milestone", "GET /health — the contract", 28)
    add_code(s, MX, 2.05, 12, 2.6, """GET http://localhost:8080/health

200 OK
{
  "status": "ok",
  "message": "Backend is running"
}""", 15)
    add_bullets(s, MX, 4.9, 12, 1.8, [
        "Keep the JSON keys exact — status and message.",
        "Also add GET /api/hello with a small JSON message (practice).",
        "Listen on :8080.",
    ], 17)
    add_notes(s, "Target on the board. Health can live in main today; routes package is fine if they are ready.")

    # 12 — Live wire health in cmd/server
    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "SMS", "Wire health into cmd/server", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 4.2, """package main
import (
    "net/http"
    "github.com/gin-gonic/gin"
)
func main() {
    r := gin.Default()
    r.GET("/health", func(c *gin.Context) {
        c.JSON(http.StatusOK, gin.H{
            "status": "ok", "message": "Backend is running",
        })
    })
    r.Run(":8080")
}""", 14)
    add_notes(s, "Replace sms starting print. go run ./cmd/server. Leave process running for curl/Postman.")

    # 13 — GET /api/hello
    s = new_slide(prs, P())
    add_kicker_title(s, "Also today", "GET /api/hello")
    add_code(s, MX, 2.05, 12, 2.2, """r.GET("/api/hello", func(c *gin.Context) {
    c.JSON(http.StatusOK, gin.H{
        "message": "hello from SMS",
    })
})""", 15)
    add_bullets(s, MX, 4.5, 12, 2.2, [
        "Groups come later; a full path is fine for day 1.",
        "Stretch homework: hard-coded GET /api/students JSON array — still no DB.",
    ], 18)
    add_notes(s, "Optional in lab if time; required-ish for stronger students as stretch.")

    # 14 — Status codes cards
    s = new_slide(prs, P())
    add_kicker_title(s, "Status codes", "Pick the number that matches the story")
    add_card(s, MX, 2.1, 5.9, 1.8, "200", "Success with a body (health, hello).")
    add_card(s, 7.0, 2.1, 5.7, 1.8, "201", "Created — next weeks for POST student.")
    add_card(s, MX, 4.1, 5.9, 1.8, "400", "Client sent junk. Your validation said no.")
    add_card(s, 7.0, 4.1, 5.7, 1.8, "404 / 500", "Missing resource / you crashed or failed internally.")
    add_notes(s, "Today mostly 200. Plant 400 for validation slide.")

    # 15 — Validation 400 light
    s = new_slide(prs, P())
    add_kicker_title(s, "Validation · light", "Reject bad input with 400", 28)
    add_code(s, MX, 2.05, 12, 2.0, """// idea for later POSTs — not required on /health
if name == "" {
    c.JSON(http.StatusBadRequest, gin.H{"error": "name required"})
    return
}""", 15)
    add_bullets(s, MX, 4.3, 12, 2.4, [
        "Health has no body — still teach the pattern once.",
        "Consistent error JSON helps Postman and future frontend.",
        "Do not panic in a handler. Return a status and a message.",
    ], 18)
    add_notes(s, "One minute. Full validation week 7/9.")

    # 16 — Live curl
    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "curl", "Prove it from the terminal", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 1.6, """curl -i http://localhost:8080/health
curl -s http://localhost:8080/api/hello | jq .""", 16)
    add_bullets(s, MX, 4.1, 12, 2.6, [
        "-i shows status line and headers.",
        "If the port is busy, someone else’s server is still running — kill it or change port temporarily.",
    ], 18)
    add_notes(s, "Students type curl themselves. jq optional.")

    # 17 — Postman
    s = new_slide(prs, P())
    add_kicker_title(s, "Postman", "The GUI you will live in")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "New request → method GET → URL http://localhost:8080/health → Send.",
        "Expect 200 and the JSON body. Screenshot for homework.",
        "Start a collection: SMS → folder Week 5 → requests Health, Hello.",
        "Save the collection; you will add Students in week 7.",
    ], 20)
    add_notes(s, "Walk the room installing Postman. WSL: localhost usually works from Windows Postman to Linux server.")

    # 18 — CORS mention only
    s = new_slide(prs, P())
    add_kicker_title(s, "CORS · mention only", "Browsers are pickier than Postman", 28)
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Postman and curl are not browsers. They ignore CORS.",
        "The React app on :5173 will need CORS later (week 15).",
        "Do not implement CORS today. Know the word exists.",
    ], 20)
    add_notes(s, "Pace valve. One slide. Move on.")

    # 19 — Where code lives
    s = new_slide(prs, P())
    add_kicker_title(s, "Project shape", "Where does this code live?")
    add_bullets(s, MX, 2.15, 12, 3.0, [
        "Minimum today: routes in cmd/server/main.go so it runs.",
        "Better: internal/routes registers /health; main only builds the engine and Run.",
        "Handlers package can wait until week 6 refactor — do not block on perfect layers.",
    ], 18)
    add_code(s, MX, 5.3, 12, 1.2, """go run ./cmd/server
# listen :8080""", 16)
    add_notes(s, "Accept messy main if health works. Praise anyone who extracts routes.")

    # 20 — Lab gate
    s = new_slide(prs, P())
    add_kicker_title(s, "Lab · rest of class", "Done when all of this is true")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "go get Gin; go.mod shows the require",
        "GET /health returns the exact JSON contract",
        "GET /api/hello returns JSON",
        "curl -i shows 200; Postman collection started",
        "Commit + push; show me the Postman 200 screenshot path or PR",
    ], 18)
    add_notes(s, "Gate on board. Typical: wrong port, forgot go get, binding 0.0.0.0 vs localhost confusion on WSL.")

    # 21 — Homework
    s = new_slide(prs, P())
    add_kicker_title(s, "Homework", "PR + Postman proof")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "PR: health + hello working on :8080.",
        "README: how to run + example curl.",
        "Attach / commit a Postman 200 screenshot (or export collection).",
        "Stretch: GET /api/students returns a hard-coded JSON array (no DB).",
    ], 20)
    add_notes(s, "Homework is the running API proof. No JWT.")

    # 22 — Next week 6
    s = new_slide(prs, P())
    add_kicker_title(s, "Next session", "Week 6 — functions and methods inside SMS")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Extract handlers from main into a handlers package.",
        "defer, multiple returns — deepen, still in your API.",
        "Bring a green health endpoint.",
    ], 20)
    add_notes(s, "Do not start student CRUD today.")

    # 23 — Words to keep cards
    s = new_slide(prs, P())
    add_kicker_title(s, "Words to keep", "Say these correctly next week")
    add_card(s, MX, 2.1, 5.9, 1.8, "Handler", "Function the router calls for a method+path.")
    add_card(s, 7.0, 2.1, 5.7, 1.8, "Status code", "Number telling the client how it went.")
    add_card(s, MX, 4.1, 5.9, 1.8, "JSON", "Text body shape for APIs.")
    add_card(s, 7.0, 4.1, 5.7, 1.8, "Gin", "Router/helpers over net/http — not magic.")
    add_notes(s, "Quick oral quiz.")

    # 24 — Safety
    s = new_slide(prs, P())
    add_kicker_title(s, "Safety", "Still true on an API week")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "No secrets in git. .env stays ignored.",
        "Do not commit the binary from go build.",
        "Port 8080 conflict: find and stop the old process.",
        "You are not graded on React today — Postman is enough.",
    ], 20)
    add_notes(s, "Frontend is week 15.")

    # 25 — Questions
    s = new_slide(prs, P(), dark=True)
    add_textbox(s, MX, 3.0, 12, 1.2, "Questions", size=64, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_notes(s, "Then stdlib demo → Gin health → curl/Postman lab.")

    if page != TOTAL:
        raise SystemExit(f"expected {TOTAL} slides, built {page}")
    prs.save(str(dest))
    print("Wrote", dest)
    return dest


if __name__ == "__main__":
    build()


