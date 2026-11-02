#!/usr/bin/env python3
"""Build an editable PowerPoint of the Week 10 Day 1 deck (native text, not screenshots)."""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    req = Path(__file__).resolve().parent / "requirements.txt"
    sys.stderr.write(
        "Missing python-pptx. From the repo root run:\n"
        f"  python3 -m pip install --user -r {req}\n"
    )
    raise

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[1]
LOGO = REPO / "docs" / "branding" / "epic-learn-logo.png"
OUT = ROOT / "pdf" / "week-10-day-01-slides.pptx"
FOOTER = "Epic Learn Institute of Higher Education - Go Programming Master Course"

INK = RGBColor(0x14, 0x20, 0x2B)
MUTED = RGBColor(0x4D, 0x5D, 0x6B)
TEAL = RGBColor(0x0B, 0x6E, 0x68)
TEAL_DARK = RGBColor(0x08, 0x4F, 0x4B)
NAVY = RGBColor(0x0A, 0x2A, 0x5C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF4, 0xEF, 0xE6)
CARD = RGBColor(0xFF, 0xFD, 0xF9)
LINE = RGBColor(0xD8, 0xCF, 0xC2)
CODE_BG = RGBColor(0x10, 0x21, 0x2C)
CODE_FG = RGBColor(0xE8, 0xF2, 0xEF)
GOLD = RGBColor(0xD9, 0xB8, 0x26)
TITLE_BG = RGBColor(0x14, 0x2A, 0x32)
CREAM_TEXT = RGBColor(0xEA, 0xD9, 0xA3)
LIVE = RGBColor(0xC2, 0x4E, 0x1D)

W, H = 13.333, 7.5
MX = 0.65
HEADER_H = 0.78
FOOTER_H = 0.42
CONTENT_TOP = 0.95
TOTAL = 26


def _font(run, name="Calibri", size=18, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_textbox(slide, l, t, w, h, text, size=18, bold=False, color=INK, font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _font(run, font, size, bold, color)
    return box


def add_bullets(slide, l, t, w, h, items, size=20, color=INK):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        p.level = 0
        run = p.add_run()
        run.text = "•  " + item
        _font(run, "Calibri", size, False, color)
    return box


def add_code(slide, l, t, w, h, text, size=15):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = CODE_BG
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.1)
    lines = text.strip("\n").split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line if line else " "
        _font(run, "Consolas", size, False, CODE_FG)
    return shape


def add_card(slide, l, t, w, h, title, body, title_size=16, body_size=13):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = LINE
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    _font(run, "Calibri", title_size, True, TEAL_DARK)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    run2 = p2.add_run()
    run2.text = body
    _font(run2, "Calibri", body_size, False, INK)
    return shape


def add_brand(slide, page, dark=False):
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(HEADER_H))
    _fill(header, WHITE)
    slide.shapes.add_picture(str(LOGO), Inches(0.35), Inches(0.12), height=Inches(0.55))
    add_textbox(
        slide, 10.2, 0.22, 2.8, 0.4,
        "WEEK 10  ·  DAY 1",
        size=12, bold=True, color=CREAM_TEXT if dark else TEAL_DARK, align=PP_ALIGN.RIGHT,
    )
    gold = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(HEADER_H), Inches(W - 0.7), Inches(0.02)
    )
    _fill(gold, GOLD)

    footer_bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(H - FOOTER_H), Inches(W), Inches(FOOTER_H)
    )
    _fill(footer_bg, WHITE)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(H - FOOTER_H), Inches(W - 0.7), Inches(0.015)
    )
    _fill(line, NAVY)
    add_textbox(
        slide, 0.5, H - FOOTER_H + 0.06, 10.5, 0.3,
        FOOTER, size=11, color=NAVY, align=PP_ALIGN.LEFT,
    )
    add_textbox(
        slide, 11.3, H - FOOTER_H + 0.06, 1.6, 0.3,
        f"{page} / {TOTAL}", size=11, color=MUTED, align=PP_ALIGN.RIGHT,
    )


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_kicker_title(slide, kicker, title, title_size=32, top=None):
    top = CONTENT_TOP if top is None else top
    add_textbox(slide, MX, top, 12, 0.32, kicker.upper(), size=13, bold=True, color=TEAL)
    add_textbox(slide, MX, top + 0.28, 12, 0.7, title, size=title_size, bold=True, color=INK)


def new_slide(prs, page, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if dark:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
        _fill(bg, TITLE_BG)
    else:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(W), Inches(H))
        _fill(bg, CREAM)
    add_brand(slide, page, dark=dark)
    return slide


def build(dest: Path | None = None) -> Path:
    dest = dest or OUT
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    page = 0

    def P():
        nonlocal page
        page += 1
        return page

    s = new_slide(prs, P(), dark=True)
    add_textbox(s, MX, 2.15, 12, 0.4, "GO PROGRAMMING MASTER COURSE", size=14, bold=True, color=CREAM_TEXT)
    add_textbox(s, MX, 2.55, 12, 1.0, "Week 10, Day 1", size=48, bold=True, color=WHITE)
    add_textbox(s, MX, 3.7, 11, 1.1, "Goroutines and the scheduler — concurrency is not a free lunch in HTTP handlers.", size=22, color=RGBColor(0xE7, 0xDD, 0xD0))
    add_textbox(s, MX, 5.3, 11, 0.9, "Umesh Indrajith · BSc Eng (Hons), University of Moratuwa\nSenior Software Engineer / Lecturer", size=16, color=RGBColor(0xD7, 0xCB, 0xB8))
    add_notes(s, "Week 9 green. Today: go f(), GMP, race demo, fan-out. Channels week 11.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Since last week", "Errors mapped. Logs structured. Still one request, one stack.", 24)
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Missing student → 404 via ErrNotFound. No panics in handlers.",
        "Cold-call: what is a goroutine?",
        "Today we start concurrent work — and see a race before optimizing SMS.",
        "Channels / worker pools wait until week 11. Mutex deep-dive week 12.",
    ], 18)
    add_notes(s, "Spot-check slog + 404.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Today", "If Day 1 of week 10 works, you can")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Say concurrency vs parallelism in one sentence each.",
        "Start work with go f() and explain it is not a thread API.",
        "Sketch GMP (Goroutine / Machine / Processor) at poster level.",
        "Spot a leaking goroutine and why handlers must not fire-and-forget.",
        "Run go test -race on a broken counter and see red.",
    ], 18)
    add_notes(s, "Safer lab: fan-out three validations.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Software engineering", "Concurrency is a product decision", 26)
    add_textbox(s, MX, 2.1, 12, 1.0, "SMS create does not need a goroutine. Reach for concurrency when latency or fan-out earns it — after you can prove safety.", size=18, color=MUTED)
    add_bullets(s, MX, 3.3, 12, 3.2, [
        "Async log-after-create is tempting — and races the map (week 12).",
        "Prefer a clear sync path today; add concurrent validation as a kata.",
        "Never go work in a handler with no cancel, no wait, no error path.",
    ], 18)
    add_notes(s, "Write no naked go in handlers on the board.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Concurrency vs parallelism", "Two words. Do not mix them.", 26)
    add_card(s, MX, 2.1, 5.9, 2.2, "Concurrency", "Many tasks in progress — structure. One CPU can still interleave them.")
    add_card(s, 7.0, 2.1, 5.7, 2.2, "Parallelism", "Many tasks at the same instant — needs multiple cores / Ps.")
    add_bullets(s, MX, 4.6, 12, 1.8, [
        "Go gives you concurrency tools. Parallelism is what the scheduler may do.",
    ], 18)
    add_notes(s, "Short.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Goroutines", "go f() — cheap concurrent function call")
    add_code(s, MX, 2.05, 12, 2.8, """go say("world")
say("hello") // may print in either order

func say(s string) {
    for i := 0; i < 3; i++ { fmt.Println(s, i) }
}""", 15)
    add_bullets(s, MX, 5.1, 12, 1.6, [
        "Not a new OS thread per go — multiplexed onto a few threads.",
        "Returns immediately; you do not wait unless you arrange it.",
    ], 16)
    add_notes(s, "Live tiny example.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Scheduler · GMP poster", "Enough picture for interviews and fear", 26)
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "G — goroutine (your go f()).",
        "M — OS thread (machine).",
        "P — logical processor (scheduling context; ~GOMAXPROCS).",
        "Runnable Gs wait on queues; Ms run Gs using Ps.",
        "Remember: many Gs, few Ms.",
    ], 18)
    add_notes(s, "One poster diagram.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Leaks", "A goroutine that never ends still costs")
    add_code(s, MX, 2.05, 12, 2.2, """go func() {
    for { /* blocked forever */ }
}()
// request finished — G still alive → leak""", 16)
    add_bullets(s, MX, 4.5, 12, 2.2, [
        "Leaks grow memory and scheduler work under load.",
        "Week 11/12: channels + context cancel as exit ramps.",
    ], 18)
    add_notes(s, "Plant context cancel.")

    s = new_slide(prs, P())
    add_kicker_title(s, "HTTP handlers", "Do not go without a plan")
    add_code(s, MX, 2.05, 12, 2.6, """func (h *StudentHandler) Create(c *gin.Context) {
    // BAD: go h.repo.Create(req)
    s := h.repo.Create(req) // GOOD for SMS today
    c.JSON(201, s)
}""", 15)
    add_bullets(s, MX, 4.9, 12, 1.8, [
        "Client needs the outcome on this request.",
        "Background work needs ownership, cancellation, and error reporting.",
    ], 18)
    add_notes(s, "Show anti-pattern then delete it.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Data race · preview", "Two goroutines, one counter, no lock", 26)
    add_code(s, MX, 2.05, 12, 2.8, """var n int
var wg sync.WaitGroup
for i := 0; i < 1000; i++ {
    wg.Add(1)
    go func() { defer wg.Done(); n++ }()
}
wg.Wait() // often n < 1000""", 14)
    add_bullets(s, MX, 5.1, 12, 1.6, [
        "go test -race will scream. Mutex fix is week 12.",
    ], 18)
    add_notes(s, "WaitGroup as glue only.")

    s = new_slide(prs, P())
    add_kicker_title(s, "go test -race", "The detector is your friend")
    add_code(s, MX, 2.05, 12, 1.8, """go test -race ./...
# WARNING: DATA RACE""", 18)
    add_bullets(s, MX, 4.1, 12, 2.6, [
        "Lab starts red on purpose.",
        "Homework: optional race on unlocked global map — document; fix next week.",
    ], 18)
    add_notes(s, "Run detector live once.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Safer lab", "Fan-out three validations")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "email / course / name checks in parallel; collect errors.",
        "Independent checks can fan out.",
        "SMS fields are tiny — sync validation is usually enough.",
        "Learn the pattern; do not force it into every handler.",
    ], 18)
    add_notes(s, "Starter has validate package + test.")

    s = new_slide(prs, P())
    add_kicker_title(s, "SMS milestone", "Practical, not theatre")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Keep create/list sync and correct.",
        "Optional later: timeout-friendly slow path / dashboard aggregation.",
        "Or: async log only after you have seen the race (week 12 fix).",
        "No channels API design this week.",
    ], 18)
    add_notes(s, "Do not merge fire-and-forget creates.")

    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "10k goroutines", "Printing is fine · shared counter is not", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 3.2, """for i := 0; i < 10_000; i++ {
    go func(i int) { fmt.Println(i) }(i)
}
// vs incrementing one int from many Gs → race""", 16)
    add_notes(s, "Capture i correctly.")

    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "-race", "Red is a lesson", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 2.4, """cd goroutine-lab
go test -race ./...
# DATA RACE on counter_test.go""", 16)
    add_bullets(s, MX, 4.9, 12, 1.8, [
        "Leave it red in class. Mutex arrives week 12.",
    ], 18)
    add_notes(s, "Students must run -race themselves.")

    s = new_slide(prs, P())
    add_textbox(s, MX, CONTENT_TOP, 3, 0.35, "LIVE DEMO", size=12, bold=True, color=LIVE)
    add_kicker_title(s, "Fan-out", "Three checks, one result", top=CONTENT_TOP + 0.28)
    add_code(s, MX, 2.2, 12, 2.4, """errs := ValidateStudentConcurrent(req)
// non-nil errs → 400 with joined messages""", 16)
    add_notes(s, "Walk starter validate package.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Anti-pattern", "Naked go in Create")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Client gets 201 before work finishes — or never learns failure.",
        "Map writes race with List/Get.",
        "Process exit can kill work mid-flight.",
        "Course rule this week: no naked go in handlers.",
    ], 20)
    add_notes(s, "Quiz them aloud.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Lab · rest of class", "Done when all of this is true")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Ran 10k-goroutine print demo (or equivalent)",
        "go test -race shows DATA RACE on the counter lab",
        "Fan-out validation runs; can explain sync vs concurrent tradeoff",
        "SMS create path still sync; health green",
        "Quiz answers started for homework",
    ], 18)
    add_notes(s, "Gate on board.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Homework", "Quiz + optional race note")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Short quiz: concurrency vs parallelism; go f(); naked go in handlers.",
        "Optional: capture -race on unsafe global map in the PR.",
        "Do not require mutex fix unless already known — week 12.",
    ], 18)
    add_notes(s, "Quiz sheet in homework.md.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Next session", "Week 11 — channels, pipelines, worker pool", 26)
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Unbuffered vs buffered; select; range/close.",
        "Worker pool as the pattern to remember.",
        "Bring race literacy and respect for naked go.",
    ], 20)
    add_notes(s, "No channel deep dive today.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Words to keep", "Say these correctly next week")
    add_card(s, MX, 2.1, 5.9, 1.8, "Goroutine", "Concurrent function scheduled by Go.")
    add_card(s, 7.0, 2.1, 5.7, 1.8, "GMP", "G=goroutine, M=thread, P=processor.")
    add_card(s, MX, 4.1, 5.9, 1.8, "Data race", "Concurrent write/read without sync.")
    add_card(s, 7.0, 4.1, 5.7, 1.8, "Leak", "Goroutine that never exits.")
    add_notes(s, "Oral quiz.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Layout reminder", "Where lab code lives")
    add_code(s, MX, 2.05, 12, 3.6, """docs/week-10-day-01/starter/
  goroutine-lab/          # race counter + fan-out validate
  backend/                # SMS still sync create/list""", 16)
    add_notes(s, "Kata separate from product path.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Honest limits", "What this week is not")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Not channels / worker pools (week 11).",
        "Not mutex mastery (week 12) — we only observe races.",
        "Not make SMS faster with go theatre.",
        "Not GOMAXPROCS tuning class.",
    ], 18)
    add_notes(s, "Keep spiral honest.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Safety", "Still true under concurrency demos")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Do not commit secrets while experimenting.",
        "Cap demo loops — 10k prints is enough.",
        "gofmt; no binaries in git.",
    ], 20)
    add_notes(s, "Quick.")

    s = new_slide(prs, P())
    add_kicker_title(s, "Review bar", "Instructor merges when")
    add_bullets(s, MX, 2.15, 12, 4.5, [
        "Quiz answers show clear concurrency vs parallelism",
        "No naked go merged into Create without discussion",
        "Race lab evidence (-race output) if they claim a race",
        "SMS health + create still green",
    ], 18)
    add_notes(s, "Collect quiz PRs.")

    s = new_slide(prs, P(), dark=True)
    add_textbox(s, MX, 3.0, 12, 1.2, "Questions", size=64, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_notes(s, "Then 10k demo → -race red → fan-out lab.")

    if page != TOTAL:
        raise SystemExit(f"expected {TOTAL} slides, built {page}")
    prs.save(str(dest))
    print("Wrote", dest)
    return dest


if __name__ == "__main__":
    build()
